"""功能切换、异常恢复与业务边界回归；全部使用临时数据库和模拟相机。"""

import json
import queue
import tempfile
import tkinter as tk
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from pathlib import Path
from tkinter import ttk
from unittest.mock import Mock, patch

import numpy as np
from openpyxl import Workbook, load_workbook

from admin_app import AdminApp, CaptureDialog
from attendance_store import AttendanceStore, read_student_import
from camera_service import CameraWorker
from GUI import Launcher
from face_utils import save_gray
from resource_lock import ResourceLease
from student_app import StudentApp
from train import manifest_path, model_health, train_model
from ui_common import INK, WHITE
from task_runtime import TaskControl


class DeferredCamera:
    def __init__(self, store, mode, **kwargs):
        self.session_id = kwargs.get("session_id")
        self.events = queue.Queue()
        self.stopped = False

    def start(self):
        pass

    def stop(self):
        self.stopped = True

    def take_frame(self):
        return None


class TemporaryStoreTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="attendance-audit-")
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.store = AttendanceStore(self.folder / "test.db", seed_path=None)
        self.store.save_student("001", "测试甲", "一班")
        self.store.save_student("002", "测试乙", "一班")
        self.now = datetime.now().replace(microsecond=0)

    def session(self, course="测试课", start=None, end=None, grace=10):
        return self.store.create_session(course, "一班", start or self.now-timedelta(minutes=2),
                                         end or self.now+timedelta(hours=1), grace)


class BusinessRegressionTests(TemporaryStoreTest):
    def test_concurrent_session_opening_only_accepts_one(self):
        sessions = [self.session("课程甲"), self.session("课程乙")]
        def open_one(sid):
            try:
                self.store.open_session(sid, self.now)
                return True
            except ValueError:
                return False
        with ThreadPoolExecutor(max_workers=2) as executor:
            self.assertEqual(sum(executor.map(open_one, sessions)), 1)

    def test_late_threshold_is_inclusive_and_end_is_exclusive(self):
        start = self.now-timedelta(minutes=10)
        sid = self.session(start=start)
        self.store.open_session(sid, self.now)
        self.assertEqual(self.store.record_sign_in("001", sid, self.now)["status"], "late")
        with self.assertRaises(ValueError):
            self.store.record_sign_in("002", sid, self.now+timedelta(hours=1))

    def test_zero_grace_marks_signin_at_start_late(self):
        sid = self.session(start=self.now, grace=0)
        self.store.open_session(sid, self.now)
        self.assertEqual(self.store.record_sign_in("001", sid, self.now)["status"], "late")

    def test_correction_audits_displayed_absence(self):
        sid = self.session()
        self.store.open_session(sid, self.now)
        self.store.close_session(sid)
        self.assertEqual(self.store.session_rows(sid)[0]["status"], "absent")
        self.store.correct_attendance(sid, "001", "present", "老师核实到课", time_unconfirmed=True)
        audit = next(row for row in self.store.audit_rows(sid) if row["action"] == "人工更正")
        self.assertIn("absent → present", audit["detail"])

    def test_manual_leave_and_absence_survive_face_checkin(self):
        sid = self.session()
        self.store.open_session(sid, self.now)
        for student, status in [("001", "leave"), ("002", "absent")]:
            self.store.correct_attendance(sid, student, status, "老师登记")
            result = self.store.record_sign_in(student, sid, self.now)
            self.assertFalse(result["inserted"])
            self.assertEqual(result["status"], status)
            self.assertIsNone(result["signed_at"])

    def test_time_validation_and_cross_day_session(self):
        for start, end in [(self.now, self.now), (self.now, self.now-timedelta(minutes=1)),
                           (self.now, self.now+timedelta(hours=25))]:
            with self.assertRaises(ValueError):
                self.session(start=start, end=end)
        start = self.now.replace(hour=23, minute=30)
        sid = self.session(start=start, end=start+timedelta(hours=2))
        self.store.open_session(sid, start)
        self.assertTrue(self.store.record_sign_in("001", sid, start)["inserted"])

    def test_expired_open_session_allows_next_session_and_audits_once(self):
        sid = self.session(end=self.now+timedelta(seconds=1))
        self.store.open_session(sid, self.now)
        future = self.now+timedelta(seconds=2)
        other = self.session("下一课")
        self.store.open_session(other, future)
        self.store.active_session(future)
        self.assertEqual(self.store.session(sid)["state"], "closed")
        self.assertEqual(sum(row["action"] == "课次到时关闭" for row in self.store.audit_rows(sid)), 1)

    def test_xlsx_import_text_ids_optional_group_and_numeric_precision(self):
        output = self.folder / "名单.xlsx"
        book = Workbook()
        book.active.append(["姓名", "学号"])
        book.active.append(["测试丙", "00000000000000000003"])
        book.save(output)
        self.assertEqual(read_student_import(output), [("00000000000000000003", "测试丙", "未分班")])
        book.active.append(["错误学生", 1234567890123456])
        book.save(output)
        book.close()
        with self.assertRaisesRegex(ValueError, "精度"):
            read_student_import(output)

    def test_export_course_and_notes_are_literal_text(self):
        sid = self.session("=1+1")
        self.store.open_session(sid, self.now)
        self.store.correct_attendance(sid, "001", "present", "=1+1", time_unconfirmed=True, note="=1+1")
        output = self.folder / "导出.xlsx"
        self.store.export_excel(output, sid)
        book = load_workbook(output)
        try:
            self.assertEqual(book.active["A1"].data_type, "s")
            self.assertEqual(book.active["H5"].data_type, "s")
        finally:
            book.close()


class GuiRegressionTests(TemporaryStoreTest):
    def setUp(self):
        super().setUp()
        self.root = tk.Tk()
        self.root.withdraw()
        self.callback_errors = []
        self.root.report_callback_exception = lambda *error: self.callback_errors.append(error)
        self.addCleanup(self.close_root)

    def close_root(self):
        try:
            self.root.update()
            for token in self.root.tk.call("after", "info"):
                self.root.after_cancel(token)
            self.root.update_idletasks()
            self.root.destroy()
        except tk.TclError:
            pass

    def student_app(self):
        sid = self.session()
        self.store.open_session(sid, self.now)
        with patch("student_app.model_health", return_value=(True, "就绪", {})):
            app = StudentApp(self.root, self.store, worker_factory=DeferredCamera, async_io=False)
            app.toggle_camera()
        return app, sid

    def test_admin_repeated_page_switches_and_refreshes_keep_selection(self):
        sid = self.session()
        app = AdminApp(self.root, self.store, authenticated=True, async_io=False)
        app.students_tree.selection_set("001")
        app.sessions_tree.selection_set(str(sid))
        self.root.update()
        for _ in range(25):
            for page in app.pages.values():
                app.notebook.select(page)
                app.refresh_students()
                app.refresh_sessions()
                self.root.update()
        self.assertFalse(self.callback_errors)
        self.assertEqual(app.students_tree.selection(), ("001",))
        self.assertEqual(app.selected_student["student_id"], "001")
        self.assertEqual(app.sessions_tree.selection(), (str(sid),))

    def test_focused_course_choice_keeps_readable_text(self):
        app = AdminApp(self.root, self.store, authenticated=True, async_io=False)
        style = ttk.Style(self.root)
        self.assertEqual(style.lookup("TCombobox", "foreground", ("readonly", "focus")), INK)
        self.assertEqual(style.lookup("TCombobox", "fieldbackground", ("readonly", "focus")), WHITE)

    def test_student_drops_previous_session_receipt_during_switch(self):
        app, sid = self.student_app()
        worker = app.worker
        record = self.store.record_sign_in("001", sid, self.now)
        self.store.close_session(sid)
        other = self.session("下一课")
        self.store.open_session(other, self.now)
        with patch("student_app.model_health", return_value=(True, "就绪", {})):
            app.refresh_session()
            worker.events.put(("signed", record))
            worker.events.put(("hint", "上一课次的识别提示"))
            app.poll_camera()
        self.assertTrue(worker.stopped)
        self.assertEqual(app.current_session["id"], other)
        self.assertEqual(app.receipt_name.cget("text"), "尚未识别")
        self.assertNotIn("上一课次", app.hint.get())

    def test_camera_stop_error_and_restart_recover_controls(self):
        app, _ = self.student_app()
        worker = app.worker
        app.toggle_camera()
        self.assertEqual(app.start_button["state"], "disabled")
        worker.events.put(("error", "模拟摄像头断开"))
        worker.events.put(("stopped", 0))
        app.poll_camera()
        self.assertIsNone(app.worker)
        self.assertEqual(app.start_button["state"], "normal")
        self.assertIn("断开", app.hint.get())
        with patch("student_app.model_health", return_value=(True, "就绪", {})):
            app.toggle_camera()
        self.assertIsNotNone(app.worker)
        self.assertIsNot(app.worker, worker)

    def test_student_shows_manual_leave_and_masks_short_id(self):
        app, sid = self.student_app()
        self.store.correct_attendance(sid, "001", "leave", "已请假")
        record = self.store.record_sign_in("001", sid, self.now)
        app.worker.events.put(("signed", record))
        app.poll_camera()
        self.assertIn("请假", app.receipt_detail.cget("text"))
        self.assertNotIn("001", app.receipt_detail.cget("text"))

    def test_model_failure_blocks_start_without_losing_page(self):
        sid = self.session()
        self.store.open_session(sid, self.now)
        with patch("student_app.model_health", return_value=(False, "样本已更新", None)):
            app = StudentApp(self.root, self.store, worker_factory=DeferredCamera, async_io=False)
            app.toggle_camera()
        self.assertIsNone(app.worker)
        self.assertIn("模型尚未就绪", app.hint.get())
        self.root.update()
        self.assertFalse(self.callback_errors)

    def test_background_completion_error_does_not_break_polling(self):
        app = AdminApp(self.root, self.store, authenticated=True, async_io=False)
        app.job_control=TaskControl()
        app.jobs.put((True, "测试任务", "结果", Mock(side_effect=ValueError("刷新失败")),app.job_control))
        with patch("ui_common.messagebox.showerror") as error:
            app.poll_jobs()
        self.assertFalse(app.busy)
        error.assert_called_once()

    def test_template_write_failure_does_not_claim_success(self):
        app = AdminApp(self.root, self.store, authenticated=True, async_io=False)
        with patch("admin_app.filedialog.asksaveasfilename", return_value=str(self.folder/"模板.csv")), \
                patch("pathlib.Path.write_text", side_effect=PermissionError("文件不可写")), \
                patch("admin_app.messagebox.showerror") as error:
            app.roster_template()
        error.assert_called_once()
        self.assertNotIn("模板已保存", app.status.get())

    def test_launcher_failure_keeps_entry_available(self):
        launcher = Launcher(self.root)
        with patch("GUI.launch_script", side_effect=OSError("启动失败")), \
                patch("GUI.messagebox.showerror") as error:
            launcher.open_role("student_app.py")
        error.assert_called_once()
        self.assertTrue(self.root.winfo_exists())

    def test_admin_and_student_follow_same_session_and_correction(self):
        admin = AdminApp(self.root, self.store, authenticated=True, async_io=False)
        student_root = tk.Toplevel(self.root)
        with patch("student_app.model_health", return_value=(True, "就绪", {})):
            student = StudentApp(student_root, self.store, worker_factory=DeferredCamera, async_io=False)
            self.assertIsNone(student.current_session)
            admin.course.set("双端联动课")
            admin.starts.set((self.now-timedelta(minutes=1)).strftime("%Y-%m-%d %H:%M"))
            admin.ends.set((self.now+timedelta(hours=1)).strftime("%Y-%m-%d %H:%M"))
            admin.create_session()
            sid = admin.selected_session_id()
            with patch("admin_app.model_health", return_value=(True, "就绪", {"student_ids": ["001", "002"]})):
                admin.open_session()
            student.refresh_session()
            self.assertEqual(student.current_session["id"], sid)
            student.toggle_camera()
            record = self.store.record_sign_in("001", sid)
            student.worker.events.put(("signed", record))
            student.poll_camera()
            admin.refresh_sessions()
            self.assertIn("已到 1", admin.record_summary.cget("text"))
            self.store.correct_attendance(sid, "002", "leave", "老师核实")
            admin.refresh_sessions()
            self.assertIn("请假 1", admin.record_summary.cget("text"))
            with patch("admin_app.messagebox.askyesno", return_value=True):
                admin.close_session()
            student.refresh_session()
            self.assertIsNone(student.current_session)
            self.assertTrue(student.worker.stopped)
            student.worker.events.put(("stopped", 0))
            student.poll_camera()
            self.assertEqual(student.start_button["state"], "disabled")
        student.close()

    def test_capture_dialog_completion_and_early_stop(self):
        completed = []
        student = self.store.students()[0]
        with patch("admin_app.CameraWorker", DeferredCamera):
            dialog = CaptureDialog(self.root, self.store, student, completed.append)
        dialog.worker.events.put(("progress", (3, 20)))
        dialog.poll()
        self.assertEqual(dialog.progress["value"], 3)
        dialog.stop()
        self.assertTrue(dialog.worker.stopped)
        dialog.worker.events.put(("stopped", 3))
        dialog.poll()
        self.assertEqual(completed, [3])
        self.assertFalse(dialog.window.winfo_exists())


class ModelRegressionTests(unittest.TestCase):
    def test_malformed_manifest_is_rejected_without_exception(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            model = folder / "model.yml"
            model.write_text("not a model", encoding="utf-8")
            for content in ["[]", "null", "42", json.dumps({"version": 2})]:
                with self.subTest(content=content):
                    manifest_path(model).write_text(content, encoding="utf-8")
                    self.assertFalse(model_health(set(), model, folder)[0])

    def test_missing_identity_list_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            sample = folder / "001.0.jpg"
            save_gray(sample,np.random.default_rng(1).integers(50,200,size=(160,160),dtype=np.uint8))
            model = folder / "model.yml"
            manifest = train_model(folder, model, {"001"})
            del manifest["student_ids"]
            manifest_path(model).write_text(json.dumps(manifest), encoding="utf-8")
            self.assertFalse(model_health({"001"}, model, folder)[0])


class CameraRegressionTests(TemporaryStoreTest):
    def test_continuous_signin_model_change_and_session_close(self):
        for transition in ("model", "samples", "session"):
            with self.subTest(transition=transition):
                sid = self.session("连续识别"+transition)
                self.store.open_session(sid, self.now)
                frame = np.random.default_rng(2).integers(50, 200, size=(240, 320, 3), dtype=np.uint8)
                capture, detector, matcher = Mock(), Mock(), Mock()
                capture.read.return_value = (True, frame)
                detector.detectMultiScale.return_value = [(30, 30, 160, 160)]
                matcher.predict.return_value = (0, 10)
                matcher.getLabelInfo.return_value = "001"
                worker = CameraWorker(self.store, "sign", session_id=sid)
                real_record = self.store.record_sign_in
                def record_and_change(*args):
                    result = real_record(*args)
                    if transition == "session":
                        self.store.close_session(sid)
                    return result
                calls = []
                def health(_ids):
                    calls.append(True)
                    if worker.events.qsize() and any(event[0] == "signed" for event in worker.events.queue):
                        if transition == "samples":
                            return False, "样本已更新，请重新训练模型", None
                        if transition == "model":
                            return True, "就绪", {"model_sha256": "new"}
                    return True, "就绪", {"model_sha256": "old"}
                # 加快时间与等待，仍走真实的连续确认、去重和资源释放路径。
                counter = iter(range(1, 1000))
                with patch("camera_service.CAMERA_LOCK_PATH", self.folder/"camera.lock"), \
                        patch("camera_service.DATA_LOCK_PATH", self.folder/"data.lock"), \
                        patch("camera_service.open_camera", return_value=capture), \
                        patch("camera_service.load_classifier", return_value=detector), \
                        patch("camera_service.load_recognizer", return_value=matcher), \
                        patch("camera_service.model_health", side_effect=health), \
                        patch("camera_service.time.monotonic", side_effect=lambda: next(counter)*.2), \
                        patch.object(worker.stop_event, "wait"), \
                        patch.object(self.store, "record_sign_in", side_effect=record_and_change):
                    worker.run()
                events = list(worker.events.queue)
                self.assertEqual(sum(kind == "signed" for kind, _ in events), 1)
                if transition == "session":
                    self.assertTrue(any(kind == "session_closed" for kind, _ in events))
                else:
                    self.assertTrue(any(kind == "error" for kind, _ in events))
                self.assertEqual(events[-1][0], "stopped")
                capture.release.assert_called_once()
                self.assertEqual(len([r for r in self.store.session_rows(sid) if r["signed_at"]]), 1)
                self.store.close_session(sid)

    def test_capture_completion_appends_samples_and_releases_camera(self):
        dataset = self.folder / "dataset"
        dataset.mkdir()
        (dataset / "old.001.7.jpg").write_bytes(b"original")
        frame = np.random.default_rng(1).integers(50, 200, size=(240, 320, 3), dtype=np.uint8)
        capture = Mock()
        capture.read.return_value = (True, frame)
        detector = Mock()
        detector.detectMultiScale.return_value = [(30, 30, 160, 160)]
        worker = CameraWorker(self.store, "capture", student_id="001", target=1)
        with patch("camera_service.DATASET_DIR", dataset), \
                patch("camera_service.sample_paths", lambda: sorted(dataset.glob("*.jpg"))), \
                patch("camera_service.CAMERA_LOCK_PATH", self.folder/"camera.lock"), \
                patch("camera_service.DATA_LOCK_PATH", self.folder/"data.lock"), \
                patch("camera_service.open_camera", return_value=capture), \
                patch("camera_service.load_classifier", return_value=detector):
            worker.run()
        self.assertEqual(worker.saved, 1)
        self.assertTrue((dataset / "001.8.jpg").exists())
        self.assertEqual((dataset / "old.001.7.jpg").read_bytes(), b"original")
        capture.release.assert_called_once()
        self.assertEqual(list(worker.events.queue)[-1], ("stopped", 1))

    def test_read_failure_releases_device_and_resource_lock(self):
        sid = self.session()
        self.store.open_session(sid, self.now)
        capture = Mock()
        capture.read.return_value = (False, None)
        worker = CameraWorker(self.store, "sign", session_id=sid)
        lock = self.folder / "camera.lock"
        with patch("camera_service.CAMERA_LOCK_PATH", lock), \
                patch("camera_service.model_health", return_value=(True, "就绪", {})), \
                patch("camera_service.load_recognizer"), patch("camera_service.load_classifier"), \
                patch("camera_service.open_camera", return_value=capture):
            worker.run()
        capture.release.assert_called_once()
        self.assertTrue(any(event[0] == "error" for event in worker.events.queue))
        self.assertEqual(list(worker.events.queue)[-1][0], "stopped")
        with ResourceLease(lock):
            pass

    def test_stopped_worker_never_opens_camera(self):
        worker = CameraWorker(self.store, "capture", student_id="001")
        worker.stop()
        with patch("camera_service.CAMERA_LOCK_PATH", self.folder/"camera.lock"), \
                patch("camera_service.DATA_LOCK_PATH", self.folder/"data.lock"), \
                patch("camera_service.open_camera") as opened:
            worker.run()
        opened.assert_not_called()


if __name__ == "__main__":
    unittest.main()
