"""课次生命周期、原始证据、更正事务与名单调整验收。"""

import json
import sqlite3
import tempfile
import tkinter as tk
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from pathlib import Path
from tkinter import ttk
from unittest.mock import patch

from openpyxl import load_workbook

from admin_app import AdminApp
from attendance_store import AttendanceStore
from _face_test_support import authorize_faces


class BusinessFixture(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="business-rules-")
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.store = AttendanceStore(self.folder / "test.db", seed_path=None)
        self.now = datetime.now().replace(microsecond=0)
        for sid, name in (("001", "甲"), ("002", "乙")):
            self.store.save_student(sid, name, "一班")
        self.start = self.now-timedelta(minutes=10)
        authorize_faces(self.store)
        self.end = self.now+timedelta(hours=1)

    def session(self, opened=False, **kwargs):
        sid = self.store.create_session("业务验收", "一班", kwargs.get("start", self.start), kwargs.get("end", self.end))
        if opened:
            self.store.open_session(sid, self.now)
        return sid

    def rows(self, sid):
        return {row["student_id"]: row for row in self.store.session_rows(sid)}

    def records(self):
        return self.store.attendance_records(self.now-timedelta(days=3), self.now+timedelta(days=3))

    def events(self, sid, action):
        return [r for r in self.store.audit_rows(sid) if r["action"] == action]

    def change_roster(self, sid, **kwargs):
        plan = self.store.preview_roster_change(sid, **kwargs)
        self.store.apply_roster_change(sid, plan["token"], "核实课次名单", actor="教师署名", **kwargs)
        return plan


class SessionStateTests(BusinessFixture):
    def test_unopened_expired_session_never_counts_absence(self):
        sid = self.session(start=self.now-timedelta(days=1), end=self.now-timedelta(hours=23))
        self.assertEqual({r["status"] for r in self.rows(sid).values()}, {"not_started"})
        self.assertEqual(self.records(), [])
        with self.assertRaisesRegex(ValueError, "只能结束"):
            self.store.close_session(sid)
        self.assertEqual(self.store.session(sid)["state"], "draft")
        self.assertEqual(self.store.sessions()[0]["expected"], 0)
        self.assertEqual(self.events(sid, "结束签到"), [])

    def test_end_is_idempotent_and_missing_session_errors(self):
        with self.assertRaisesRegex(ValueError, "不存在"):
            self.store.close_session(999999)
        sid = self.session(opened=True)
        with ThreadPoolExecutor(max_workers=4) as executor:
            results = list(executor.map(self.store.close_session, [sid]*4))
        self.assertEqual(sum(results), 1)
        self.assertEqual(len(self.events(sid, "结束签到")), 1)
        self.assertEqual({r["status"] for r in self.rows(sid).values()}, {"absent"})

    def test_cancel_preserves_roster_and_reason_but_excludes_reports(self):
        sid = self.session()
        with self.assertRaises(ValueError):
            self.store.cancel_session(sid, " ")
        self.assertTrue(self.store.cancel_session(sid, "课程临时取消", "王老师"))
        self.assertFalse(self.store.cancel_session(sid, "重复点击"))
        self.assertEqual(self.store.session(sid)["cancel_reason"], "课程临时取消")
        self.assertEqual({r["status"] for r in self.rows(sid).values()}, {"cancelled"})
        self.assertEqual(len(self.rows(sid)), 2)
        self.assertEqual(self.records(), [])
        self.assertEqual(len(self.events(sid, "取消课次")), 1)
        for call in (lambda: self.store.open_session(sid, self.now), lambda: self.store.close_session(sid),
                     lambda: self.store.correct_attendance(sid, "001", "leave", "原因"),
                     lambda: self.store.preview_roster_change(sid, refresh=True),
                     lambda: self.store.export_excel(self.folder/"cancelled.xlsx", sid)):
            with self.assertRaises(ValueError):
                call()

    def test_open_or_closed_session_cannot_be_cancelled(self):
        sid = self.session(opened=True)
        with self.assertRaises(ValueError):
            self.store.cancel_session(sid, "错误操作")
        self.store.close_session(sid)
        with self.assertRaises(ValueError):
            self.store.cancel_session(sid, "错误操作")

    def test_expiry_detail_aggregate_and_export_agree(self):
        sid = self.session(opened=True)
        future = self.end+timedelta(seconds=1)
        detail = self.store.session_rows(sid, now=future)
        summary = self.store.attendance_records(self.start, future, now=future)
        self.assertEqual([(r["student_id"], r["status"]) for r in detail], [(r["student_id"], r["status"]) for r in summary])
        self.store.session_rows(sid, now=future)
        self.assertEqual(len(self.events(sid, "课次到时关闭")), 1)
        self.assertFalse(self.store.close_session(sid))
        self.assertEqual(self.events(sid, "结束签到"), [])
        output = self.folder/"ended.xlsx"
        self.store.export_excel(output, sid)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual([book.active.cell(r, 5).value for r in (5, 6)], ["缺勤", "缺勤"])

    def test_cancel_and_audit_fail_together(self):
        sid = self.session()
        with patch.object(self.store, "_audit", side_effect=sqlite3.OperationalError("audit failed")):
            with self.assertRaises(sqlite3.Error):
                self.store.cancel_session(sid, "原因")
        self.assertEqual(self.store.session(sid)["state"], "draft")


class CorrectionEvidenceTests(BusinessFixture):
    def test_note_and_status_changes_preserve_original_arrival(self):
        sid = self.session(opened=True)
        original = self.store.record_sign_in("001", sid, self.now-timedelta(minutes=7))
        self.store.correct_attendance(sid, "001", "present", "补充说明", note="已核实", actor="王老师")
        self.assertEqual(self.rows(sid)["001"]["signed_at"], original["signed_at"])
        self.store.correct_attendance(sid, "001", "leave", "重新核对")
        self.assertEqual(self.rows(sid)["001"]["signed_at"], original["signed_at"])
        history = self.store.attendance_history(sid, "001")
        self.assertEqual(history["evidence"]["signed_at"], original["signed_at"])
        self.assertEqual(len(history["corrections"]), 2)
        first = history["corrections"][0]
        self.assertEqual((first["actor_kind"], first["actor_label"]), ("management_signature", "王老师"))
        before, after = json.loads(first["before_json"]), json.loads(first["after_json"])
        self.assertEqual((before["method"], before["signed_at"]), ("face", original["signed_at"]))
        self.assertEqual((after["method"], after["note"]), ("manual", "已核实"))
        self.assertGreater(first["created_at"], original["signed_at"])

    def test_manual_signin_requires_explicit_time_or_unknown(self):
        sid = self.session(opened=True)
        with self.assertRaisesRegex(ValueError, "实际到课时间"):
            self.store.correct_attendance(sid, "001", "present", "老师核实")
        self.assertEqual(self.events(sid, "人工更正"), [])
        self.store.correct_attendance(sid, "001", "present", "时间未知", time_unconfirmed=True)
        self.assertIsNone(self.rows(sid)["001"]["signed_at"])
        self.store.correct_attendance(sid, "002", "late", "核实迟到", arrival_time=self.now-timedelta(minutes=1))
        self.assertEqual(self.rows(sid)["002"]["signed_at"], (self.now-timedelta(minutes=1)).isoformat())
        self.assertEqual(self.store.attendance_history(sid, "001")["corrections"][0]["actor_kind"], "management_entry")

    def test_unknown_time_remains_unknown_on_later_note_change(self):
        sid = self.session(opened=True)
        self.store.correct_attendance(sid, "001", "present", "时间未知", time_unconfirmed=True)
        self.store.correct_attendance(sid, "001", "present", "补充原因", note="核实到课")
        self.assertIsNone(self.rows(sid)["001"]["signed_at"])

    def test_actual_time_override_does_not_modify_evidence(self):
        sid = self.session(opened=True)
        first = self.store.record_sign_in("001", sid, self.now-timedelta(minutes=5))
        self.store.correct_attendance(sid, "001", "late", "核实时间", arrival_time=self.now)
        self.assertEqual(self.rows(sid)["001"]["signed_at"], self.now.isoformat())
        self.assertEqual(self.rows(sid)["001"]["original_signed_at"], first["signed_at"])
        for statement in ("UPDATE face_evidence SET signed_at='wrong'", "DELETE FROM face_evidence"):
            with self.assertRaises(sqlite3.IntegrityError):
                with self.store.connect() as connection:
                    connection.execute(statement)

    def test_repeated_identical_correction_is_one_event(self):
        sid = self.session(opened=True)
        with ThreadPoolExecutor(max_workers=3) as executor:
            results = list(executor.map(lambda _: self.store.correct_attendance(sid, "001", "present", "已核实", time_unconfirmed=True), range(3)))
        self.assertEqual(sum(results), 1)
        self.assertEqual(len(self.events(sid, "人工更正")), 1)
        self.store.correct_attendance(sid, "001", "present", "新原因", time_unconfirmed=True)
        self.assertEqual(len(self.events(sid, "人工更正")), 2)

    def test_audit_failure_rolls_back_current_result(self):
        sid = self.session(opened=True)
        self.store.record_sign_in("001", sid, self.now)
        previous = self.rows(sid)["001"]
        with patch.object(self.store, "_audit", side_effect=sqlite3.OperationalError("audit failed")):
            with self.assertRaises(sqlite3.Error):
                self.store.correct_attendance(sid, "001", "absent", "核实")
        self.assertEqual(self.rows(sid)["001"], previous)
        self.assertEqual(self.events(sid, "人工更正"), [])

    def test_face_after_manual_result_only_adds_immutable_evidence(self):
        sid = self.session(opened=True)
        self.store.correct_attendance(sid, "001", "leave", "请假")
        record = self.store.record_sign_in("001", sid, self.now)
        self.assertFalse(record["inserted"])
        self.assertEqual((record["status"], record["signed_at"]), ("leave", None))
        self.store.record_sign_in("001", sid, self.now+timedelta(seconds=1))
        self.assertEqual(self.store.attendance_history(sid, "001")["evidence"]["signed_at"], self.now.isoformat())
        self.assertEqual(len(self.events(sid, "保存首次人脸证据")), 1)

    def test_excel_marks_unknown_time_and_preserves_evidence_column(self):
        sid = self.session(opened=True)
        self.store.correct_attendance(sid, "001", "present", "未知时间", time_unconfirmed=True)
        self.store.record_sign_in("002", sid, self.now)
        self.store.correct_attendance(sid, "002", "absent", "核实缺勤")
        output = self.folder/"results.xlsx"
        self.store.export_attendance(output, self.start, self.end)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active["M5"].value, "时间未确认")
        self.assertIsNone(book.active["F5"].value)
        self.assertEqual(book.active["L6"].value, self.now)
        self.assertEqual(book.active["F6"].value, self.now)


class RosterAdjustmentTests(BusinessFixture):
    def test_draft_refresh_preview_preserves_original_rows(self):
        sid = self.session()
        self.store.save_student("003", "丙", "一班")
        self.store.save_student("002", "乙新名", "二班")
        plan = self.store.preview_roster_change(sid, refresh=True)
        self.assertEqual((plan["added"], plan["removed"]), (["003"], ["002"]))
        self.assertEqual(set(self.rows(sid)), {"001", "002"})
        self.store.apply_roster_change(sid, plan["token"], "核对班级", refresh=True)
        self.assertEqual(set(self.rows(sid)), {"001", "003"})
        archived = {r["student_id"]: r for r in self.store.session_rows(sid, include_removed=True)}
        self.assertEqual((archived["002"]["name"], archived["002"]["status"]), ("乙", "excluded"))
        self.assertEqual(len(self.events(sid, "调整课次名单")), 1)

    def test_open_add_and_soft_remove_retains_checkin_and_evidence(self):
        sid = self.session(opened=True)
        self.store.record_sign_in("001", sid, self.now)
        self.store.save_student("003", "丙", "二班")
        authorize_faces(self.store)
        plan = self.change_roster(sid, add_ids=["003"], remove_ids=["001"])
        self.assertEqual(plan["recorded_removals"], ["001"])
        with self.assertRaises(ValueError):
            self.store.record_sign_in("001", sid, self.now)
        self.assertTrue(self.store.record_sign_in("003", sid, self.now)["inserted"])
        self.assertIsNotNone(self.store.attendance_history(sid, "001")["evidence"])
        with self.store.connect() as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM checkins WHERE session_id=?", (sid,)).fetchone()[0], 2)
        self.assertEqual({r["student_id"] for r in self.records()}, {"002", "003"})
        summary = next(s for s in self.store.sessions() if s["id"] == sid)
        self.assertEqual((summary["expected"], summary["arrived"]), (2, 1))

    def test_closed_change_is_audited_and_reincluded_record_survives(self):
        sid = self.session(opened=True)
        first = self.store.record_sign_in("001", sid, self.now)
        self.store.close_session(sid)
        self.change_roster(sid, remove_ids=["001"])
        self.assertEqual(len(self.records()), 1)
        self.change_roster(sid, add_ids=["001"])
        self.assertEqual(self.rows(sid)["001"]["signed_at"], first["signed_at"])
        self.assertEqual(len(self.events(sid, "调整课次名单")), 2)
        change = self.events(sid, "调整课次名单")[0]
        self.assertTrue(change["before_json"] and change["after_json"])
        self.assertEqual(change["actor_kind"], "management_signature")

    def test_stale_preview_is_rejected_after_checkin_or_state_change(self):
        sid = self.session(opened=True)
        plan = self.store.preview_roster_change(sid, remove_ids=["001"])
        self.store.record_sign_in("001", sid, self.now)
        with self.assertRaisesRegex(ValueError, "重新预览"):
            self.store.apply_roster_change(sid, plan["token"], "核实", remove_ids=["001"])
        plan = self.store.preview_roster_change(sid, remove_ids=["001"])
        self.store.close_session(sid)
        with self.assertRaisesRegex(ValueError, "重新预览"):
            self.store.apply_roster_change(sid, plan["token"], "核实", remove_ids=["001"])
        self.assertEqual(set(self.rows(sid)), {"001", "002"})

    def test_disable_does_not_change_roster_and_new_sessions_skip_disabled(self):
        sid = self.session(opened=True)
        self.store.set_student_active("002", False)
        self.assertEqual(set(self.rows(sid)), {"001", "002"})
        with self.assertRaises(ValueError):
            self.store.record_sign_in("002", sid, self.now)
        self.assertEqual(set(self.rows(self.session())), {"001"})

    def test_roster_audit_failure_rolls_back_and_noop_writes_no_event(self):
        sid = self.session()
        plan = self.store.preview_roster_change(sid, refresh=True)
        self.assertFalse(self.store.apply_roster_change(sid, plan["token"], "无变化", refresh=True))
        plan = self.store.preview_roster_change(sid, remove_ids=["002"])
        with patch.object(self.store, "_audit", side_effect=sqlite3.OperationalError("audit failed")):
            with self.assertRaises(sqlite3.Error):
                self.store.apply_roster_change(sid, plan["token"], "核实", remove_ids=["002"])
        self.assertEqual(set(self.rows(sid)), {"001", "002"})
        self.assertEqual(self.events(sid, "调整课次名单"), [])

    def test_empty_roster_remains_traceable_and_cannot_be_opened(self):
        sid = self.session()
        self.store.set_student_active("001", False)
        self.store.set_student_active("002", False)
        self.change_roster(sid, refresh=True)
        self.assertEqual(self.rows(sid), {})
        self.assertEqual(len(self.store.session_rows(sid, include_removed=True)), 2)
        with self.assertRaisesRegex(ValueError, "名单为空"):
            self.store.open_session(sid, self.now)


class MigrationTests(unittest.TestCase):
    def test_version_two_migrates_once_and_does_not_invent_lost_evidence(self):
        with tempfile.TemporaryDirectory(prefix="business-migrate-") as folder:
            path = Path(folder)/"old.db"
            with sqlite3.connect(path) as connection:
                connection.executescript("""
                    CREATE TABLE metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL);
                    INSERT INTO metadata VALUES('schema_version','2'),('seed_imported','1');
                    CREATE TABLE students(student_id TEXT PRIMARY KEY,name TEXT NOT NULL,group_name TEXT,active INTEGER DEFAULT 1);
                    INSERT INTO students VALUES('001','甲','一班',1),('002','乙','一班',1);
                    CREATE TABLE sessions(id INTEGER PRIMARY KEY,course_name TEXT,group_name TEXT,starts_at TEXT,late_at TEXT,ends_at TEXT,state TEXT,created_at TEXT);
                    INSERT INTO sessions VALUES(1,'原课程','一班','2026-10-21T08:00:00','2026-10-21T08:10:00','2026-10-21T09:00:00','closed','2026-10-21T07:00:00');
                    CREATE TABLE session_students(session_id INTEGER,student_id TEXT,name TEXT,group_name TEXT,PRIMARY KEY(session_id,student_id));
                    INSERT INTO session_students VALUES(1,'001','甲','一班'),(1,'002','乙','一班');
                    CREATE TABLE checkins(session_id INTEGER,student_id TEXT,status TEXT,signed_at TEXT,method TEXT,note TEXT,PRIMARY KEY(session_id,student_id));
                    INSERT INTO checkins VALUES(1,'001','present','2026-10-21T08:03:00','face',''),(1,'002','present','2026-10-21T10:00:00','manual','旧更正');
                    CREATE TABLE audit_log(id INTEGER PRIMARY KEY,created_at TEXT,action TEXT,session_id INTEGER,student_id TEXT,detail TEXT);
                    INSERT INTO audit_log VALUES(1,'2026-10-21T10:00:00','人工更正',1,'002','旧版原因');
                """)
            connection.close()
            store = AttendanceStore(path, seed_path=None)
            AttendanceStore(path, seed_path=None)
            self.assertEqual(store.attendance_history(1, "001")["evidence"]["signed_at"], "2026-10-21T08:03:00")
            self.assertIsNone(store.attendance_history(1, "002")["evidence"])
            old = store.attendance_history(1, "002")["corrections"][0]
            self.assertEqual(old["actor_kind"], "legacy_unknown")
            self.assertIsNone(old["before_json"])
            with store.connect() as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM sessions").fetchone()[0], 1)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM face_evidence").fetchone()[0], 1)
                self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")
                self.assertEqual(connection.execute("PRAGMA foreign_key_check").fetchall(), [])


class BusinessGuiTests(BusinessFixture):
    def setUp(self):
        super().setUp()
        self.root = tk.Tk()
        self.root.withdraw()
        self.errors = []
        self.root.report_callback_exception = lambda *args: self.errors.append(str(args))
        self.addCleanup(self.close_root)
        self.app = AdminApp(self.root, self.store, authenticated=True, async_io=False)

    def close_root(self):
        try:
            for token in self.root.tk.call("after", "info"):
                self.root.after_cancel(token)
            self.root.destroy()
        except tk.TclError:
            pass

    def descendants(self, widget):
        for child in widget.winfo_children():
            yield child
            yield from self.descendants(child)

    def dialog(self):
        return next(w for w in self.root.winfo_children() if isinstance(w, tk.Toplevel))

    def select(self, sid):
        self.app.refresh_sessions()
        self.app.sessions_tree.selection_set(str(sid))
        self.app.sync_session_actions()

    def test_action_buttons_follow_state_and_cancel_reason_is_queryable(self):
        sid = self.session()
        self.select(sid)
        self.assertEqual(self.app.close_session_button.cget("state"), "disabled")
        self.assertEqual(self.app.cancel_session_button.cget("state"), "normal")
        self.app.cancel_session()
        dialog = self.dialog()
        entries = [w for w in self.descendants(dialog) if isinstance(w, ttk.Entry)]
        entries[0].insert(0, "调课取消")
        entries[1].insert(0, "张老师")
        next(w for w in self.descendants(dialog) if isinstance(w, tk.Button) and w.cget("text") == "确认取消").invoke()
        self.assertEqual(self.store.session(sid)["state"], "cancelled")
        self.assertEqual(self.app.open_session_button.cget("state"), "disabled")
        self.app.show_session_detail()
        content = next(w for w in self.descendants(self.dialog()) if isinstance(w, tk.Text)).get("1.0", "end")
        self.assertIn("调课取消", content)
        self.assertEqual(self.errors, [])

    def test_gui_unknown_time_requires_explicit_choice(self):
        sid = self.session(opened=True)
        # 精确使用课次日期，包含零点前创建、零点后仍在进行的课次。
        self.app.record_period.set('选定课次日期')
        self.app.record_target = self.store.session(sid)
        self.select(sid)
        self.app.records_tree.selection_set('001')
        self.app.correct_record()
        dialog = self.dialog()
        next(w for w in self.descendants(dialog) if isinstance(w, tk.Text)).insert("1.0", "现场核实")
        save = next(w for w in self.descendants(dialog) if isinstance(w, tk.Button) and w.cget("text") == "保存更正")
        with patch("admin_app.messagebox.showerror") as error:
            save.invoke()
            self.assertTrue(error.called)
        next(w for w in self.descendants(dialog) if isinstance(w, ttk.Checkbutton)).invoke()
        save.invoke()
        self.assertEqual((self.rows(sid)["001"]["status"], self.rows(sid)["001"]["signed_at"]), ("present", None))
        self.assertIn("时间未确认", self.app.records_tree.item('001', "values"))
        self.app.show_selected_history()
        history = next(w for w in self.descendants(self.dialog()) if isinstance(w, tk.Text)).get("1.0", "end")
        self.assertIn("管理入口（未署名）", history)
        self.assertIn("状态：已签到", history)
        self.assertIn("状态：未签到", history)
        self.assertEqual(self.errors, [])

    def test_gui_roster_preview_does_not_apply_until_confirmed(self):
        sid = self.session()
        self.store.save_student("003", "丙", "一班")
        self.select(sid)
        self.app.manage_roster()
        dialog = self.dialog()
        self.root.update_idletasks()
        self.assertLessEqual(dialog.winfo_reqheight(), 700)
        buttons = {w.cget("text"): w for w in self.descendants(dialog) if isinstance(w, tk.Button)}
        buttons["预览班级刷新"].invoke()
        self.assertEqual(set(self.rows(sid)), {"001", "002"})
        entries = [w for w in self.descendants(dialog) if isinstance(w, ttk.Entry)]
        entries[2].insert(0, "新同学转入")
        buttons["确认保存变更"].invoke()
        self.assertEqual(set(self.rows(sid)), {"001", "002", "003"})
        self.assertEqual(self.errors, [])


if __name__ == "__main__":
    unittest.main()
