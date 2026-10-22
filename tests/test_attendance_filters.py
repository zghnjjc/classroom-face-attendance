"""日期筛选、跨课次汇总、缺勤推导及筛选后更正/导出的回归检查。"""

import sqlite3
import tempfile
import tkinter as tk
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from tkinter import ttk
from unittest.mock import patch

from openpyxl import load_workbook

from admin_app import AdminApp
from attendance_filters import default_semester_dates, record_date_range
from attendance_store import AttendanceStore


class DateRangeTests(unittest.TestCase):
    def test_day_week_and_month_boundaries(self):
        now = datetime(2026, 10, 21, 14, 30)
        self.assertEqual(record_date_range("当天", now), (datetime(2026, 10, 21), datetime(2026, 10, 22)))
        self.assertEqual(record_date_range("当周", now), (datetime(2026, 10, 19), datetime(2026, 10, 26)))
        self.assertEqual(record_date_range("当月", now), (datetime(2026, 10, 1), datetime(2026, 11, 1)))
        self.assertEqual(record_date_range("当周", datetime(2027, 1, 3)), (datetime(2026, 12, 28), datetime(2027, 1, 4)))
        self.assertEqual(record_date_range("当月", datetime(2026, 12, 31)), (datetime(2026, 12, 1), datetime(2027, 1, 1)))
        self.assertEqual(record_date_range("当月", datetime(2028, 2, 29)), (datetime(2028, 2, 1), datetime(2028, 3, 1)))

    def test_semester_year_boundary_and_custom_inclusive_end(self):
        self.assertEqual(default_semester_dates(datetime(2027, 1, 31)), (date(2026, 9, 1), date(2027, 1, 31)))
        self.assertEqual(default_semester_dates(datetime(2027, 2, 1)), (date(2027, 2, 1), date(2027, 8, 31)))
        self.assertEqual(default_semester_dates(datetime(2027, 8, 31)), (date(2027, 2, 1), date(2027, 8, 31)))
        self.assertEqual(default_semester_dates(datetime(2027, 9, 1)), (date(2027, 9, 1), date(2028, 1, 31)))
        self.assertEqual(record_date_range("当前学期", datetime(2026, 10, 21), ("2026-08-24", "2027-01-17")),
                         (datetime(2026, 8, 24), datetime(2027, 1, 18)))


class FilterFixture(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="attendance-filter-")
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.store = AttendanceStore(self.folder / "test.db", seed_path=None)
        for sid, name in (("001", "甲"), ("002", "乙"), ("003", "丙")):
            self.store.save_student(sid, name, "一班")
        self.now = datetime(2026, 10, 21, 12)

    def closed_session(self, start, name="测试课程", signed=()):
        sid = self.store.create_session(name, "一班", start, start + timedelta(hours=1))
        self.store.open_session(sid, start)
        for student in signed:
            self.store.record_sign_in(student, sid, start + timedelta(minutes=1))
        self.store.close_session(sid)
        return sid


class FilterStoreTests(FilterFixture):
    def test_time_filter_uses_session_start_with_exclusive_end(self):
        previous = self.closed_session(datetime(2026, 10, 20, 23, 30))
        first = self.closed_session(datetime(2026, 10, 21))
        second = self.closed_session(datetime(2026, 10, 21, 23, 59, 59))
        following = self.closed_session(datetime(2026, 10, 22))
        start, end = record_date_range("当天", self.now)
        rows = self.store.attendance_records(start, end, now=self.now)
        self.assertEqual({r["session_id"] for r in rows}, {first, second})
        self.assertEqual(len(rows), 6)
        self.assertNotIn(previous, {r["session_id"] for r in rows})
        self.assertNotIn(following, {r["session_id"] for r in rows})

    def test_week_month_and_semester_select_distinct_sessions(self):
        starts = [datetime(2026, 8, 31), datetime(2026, 9, 1), datetime(2026, 9, 30),
                  datetime(2026, 10, 1), datetime(2026, 10, 18), datetime(2026, 10, 19),
                  datetime(2026, 10, 25, 23, 59, 59), datetime(2026, 10, 26), datetime(2026, 11, 1)]
        ids = [self.closed_session(start) for start in starts]
        expected = {"当周": set(ids[5:7]), "当月": set(ids[3:8]), "当前学期": set(ids[1:])}
        for period, wanted in expected.items():
            with self.subTest(period=period):
                rows = self.store.attendance_records(*record_date_range(period, self.now), now=self.now)
                self.assertEqual({r["session_id"] for r in rows}, wanted)

    def test_absent_filter_excludes_pending_leave_and_signed(self):
        closed = self.closed_session(datetime(2026, 10, 21, 8), signed=("001",))
        self.store.correct_attendance(closed, "002", "leave", "已核实请假")
        pending = self.store.create_session("进行中的课程", "一班", self.now, self.now + timedelta(days=1))
        self.store.open_session(pending, self.now)
        draft = self.store.create_session("尚未开启", "一班", self.now, self.now + timedelta(hours=1))
        start, end = record_date_range("当天", self.now)
        all_rows = self.store.attendance_records(start, end, now=self.now)
        self.assertEqual(len(all_rows), 6)
        self.assertNotIn(draft, {r["session_id"] for r in all_rows})
        self.assertTrue(all(r["status"] == "pending" for r in all_rows if r["session_id"] == pending))
        absent = self.store.attendance_records(start, end, "absent", now=self.now)
        self.assertEqual([(r["session_id"], r["student_id"]) for r in absent], [(closed, "003")])
        expired = self.store.attendance_records(start, end, "absent", now=self.now + timedelta(days=2))
        self.assertEqual(len(expired), 4)

    def test_manual_absence_and_snapshot_remain_correct_after_student_edit(self):
        sid = self.closed_session(datetime(2026, 10, 21), signed=("001",))
        original_time = self.store.session_rows(sid)[0]["signed_at"]
        self.store.correct_attendance(sid, "001", "absent", "老师核实缺勤", note="老师核实缺勤")
        self.store.save_student("001", "新名字", "另一班")
        self.store.set_student_active("001", False)
        rows = self.store.attendance_records(*record_date_range("当天", self.now), "absent", now=self.now)
        row = next(r for r in rows if r["student_id"] == "001")
        self.assertEqual((row["name"], row["group_name"], row["method"], row["signed_at"]), ("甲", "一班", "manual", original_time))
        self.assertEqual(row["original_signed_at"], original_time)
        self.assertEqual(row["note"], "老师核实缺勤")

    def test_custom_semester_persists_and_reset_restores_automatic_rollover(self):
        self.store.set_semester_dates("2026-08-24", "2027-01-17")
        restarted = AttendanceStore(self.store.path, seed_path=None)
        self.assertEqual(restarted.semester_dates(self.now), (date(2026, 8, 24), date(2027, 1, 17)))
        for start, end in (("2026-02-30", "2026-06-01"), ("2026-10-02", "2026-10-01")):
            with self.assertRaises(ValueError):
                restarted.set_semester_dates(start, end)
        self.assertEqual(restarted.semester_dates(self.now), (date(2026, 8, 24), date(2027, 1, 17)))
        restarted.reset_semester_dates()
        self.assertEqual(restarted.semester_dates(datetime(2027, 2, 1)), (date(2027, 2, 1), date(2027, 8, 31)))

    def test_export_matches_absence_and_course_filters_with_context_and_text_ids(self):
        first = self.closed_session(datetime(2026, 10, 21, 8), name="=课程名称", signed=("001",))
        self.closed_session(datetime(2026, 10, 21, 9))
        self.closed_session(datetime(2026, 10, 20))
        start, end = record_date_range("当天", self.now)
        output = self.folder / "缺勤.xlsx"
        self.store.export_attendance(output, start, end, "absent", first, now=self.now)
        book = load_workbook(output)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual([sheet.cell(r, 2).value for r in (5, 6)], ["002", "003"])
        self.assertTrue(all(sheet.cell(r, 5).value == "缺勤" for r in (5, 6)))
        self.assertEqual(sheet["I5"].value, "=课程名称")
        self.assertEqual(sheet["I5"].data_type, "s")
        self.assertEqual(sheet["B5"].number_format, "@")
        self.assertEqual(sheet["J5"].value, first)
        self.assertEqual(sheet["K5"].value, datetime(2026, 10, 21, 8))
        self.assertEqual(sheet.auto_filter.ref, "A4:M6")
        self.assertIsNone(sheet["F5"].value)

    def test_empty_range_and_invalid_query_do_not_broaden_results(self):
        self.closed_session(datetime(2026, 10, 20))
        start, end = record_date_range("当天", self.now)
        self.assertEqual(self.store.attendance_records(start, end, now=self.now), [])
        with self.assertRaises(ValueError):
            self.store.attendance_records(end, start)
        with self.assertRaises(ValueError):
            self.store.attendance_records(start, end, "unknown")
        output = self.folder / "空记录.xlsx"
        self.store.export_attendance(output, start, end, now=self.now)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active.max_row, 4)


class FilterGuiTests(FilterFixture):
    def setUp(self):
        super().setUp()
        self.clock = type("Clock", (datetime,), {"current": self.now,
            "now": classmethod(lambda cls, tz=None: cls.current)})
        self.clock_patch = patch("admin_app.datetime", self.clock)
        self.clock_patch.start()
        self.addCleanup(self.clock_patch.stop)
        self.first = self.closed_session(datetime(2026, 10, 21, 8), "课程一", signed=("001",))
        self.second = self.closed_session(datetime(2026, 10, 21, 9), "课程二", signed=("002",))
        self.older = self.closed_session(datetime(2026, 10, 13), "上周课程")
        self.root = tk.Tk()
        self.root.withdraw()
        self.callback_errors = []
        self.root.report_callback_exception = lambda *args: self.callback_errors.append(str(args[1]))
        self.addCleanup(self.cleanup_root)
        self.app = AdminApp(self.root, self.store, authenticated=True, async_io=False)

    def cleanup_root(self):
        try:
            self.root.update()
            for token in self.root.tk.call("after", "info"):
                self.root.after_cancel(token)
            self.root.update_idletasks()
            self.root.destroy()
        except tk.TclError:
            pass

    def choose(self, variable, widget, value):
        variable.set(value)
        widget.event_generate("<<ComboboxSelected>>")
        self.root.update()

    def test_gui_defaults_show_all_today_with_distinct_course_rows(self):
        self.assertEqual((self.app.record_period.get(), self.app.record_status.get(), self.app.record_choice.get()), ("当天", "全部", "全部课次"))
        self.assertEqual(len(self.app.records_tree.get_children()), 6)
        self.assertEqual({r["session_id"] for r in self.app.record_rows}, {self.first, self.second})
        self.assertTrue(self.app.records_tree.exists(f"{self.first}:001"))
        self.assertTrue(self.app.records_tree.exists(f"{self.second}:001"))

    def test_switching_filters_and_periodic_refresh_keep_selected_scope(self):
        self.choose(self.app.record_period, self.app.record_period_combo, "当月")
        self.choose(self.app.record_status, self.app.record_status_combo, "缺勤")
        for _ in range(4):
            for page in self.app.pages.values():
                self.app.notebook.select(page)
                self.app.refresh_sessions()
                self.root.update()
        self.assertEqual(len(self.app.record_rows), 7)
        self.assertTrue(all(r["status"] == "absent" for r in self.app.record_rows))
        self.assertEqual((self.app.record_period.get(), self.app.record_status.get()), ("当月", "缺勤"))
        self.assertFalse(self.callback_errors)

    def test_midnight_moves_today_filter_without_showing_previous_day(self):
        choice = next(k for k, v in self.app.record_lookup.items() if v == self.first)
        self.choose(self.app.record_choice, self.app.record_combo, choice)
        self.clock.current += timedelta(days=1)
        self.app.tick()
        self.assertEqual(self.app.record_period.get(), "当天")
        self.assertEqual(self.app.record_choice.get(), "全部课次")
        self.assertEqual(self.app.record_rows, [])
        self.assertIn("2026-10-22", self.app.record_range_label.cget("text"))

    def test_period_query_failure_recovers_without_resetting_filters(self):
        self.choose(self.app.record_status, self.app.record_status_combo, "缺勤")
        with patch.object(self.store, "attendance_records", side_effect=sqlite3.OperationalError("database is locked")):
            self.app.tick()
        self.assertTrue(self.app.refresh_error)
        self.app.tick()
        self.assertFalse(self.app.refresh_error)
        self.assertEqual(self.app.record_status.get(), "缺勤")
        self.assertEqual(len(self.app.record_rows), 4)

    def test_correction_targets_selected_course_and_removes_row_from_absence_filter(self):
        self.choose(self.app.record_status, self.app.record_status_combo, "缺勤")
        self.app.records_tree.selection_set(f"{self.second}:001")
        original = self.store.session_rows(self.first)[0]
        self.app.correct_record()
        dialog = next(w for w in self.root.winfo_children() if isinstance(w, tk.Toplevel))
        combo = next(w for w in dialog.winfo_children() if isinstance(w, ttk.Combobox))
        reason = next(w for w in dialog.winfo_children() if isinstance(w, tk.Text))
        combo.set("请假"); reason.insert("1.0", "按假条更正")
        save = next(w for w in dialog.winfo_children() if isinstance(w, tk.Button) and w.cget("text") == "保存更正")
        save.invoke()
        self.assertEqual(self.store.session_rows(self.first)[0], original)
        changed = self.store.session_rows(self.second)[0]
        self.assertEqual((changed["status"], changed["note"]), ("leave", ""))
        self.assertIn("按假条更正", self.store.attendance_history(self.second,"001")["corrections"][-1]["detail"])
        self.assertNotIn(f"{self.second}:001", self.app.records_tree.get_children())

    def test_gui_export_uses_current_time_status_and_course_filters(self):
        self.choose(self.app.record_status, self.app.record_status_combo, "缺勤")
        choice = next(k for k, v in self.app.record_lookup.items() if v == self.second)
        self.choose(self.app.record_choice, self.app.record_combo, choice)
        output = self.folder / "界面筛选导出.xlsx"
        with patch("admin_app.filedialog.asksaveasfilename", return_value=str(output)):
            self.app.export()
        book = load_workbook(output)
        self.addCleanup(book.close)
        exported = list(book.active.iter_rows(min_row=5, values_only=True))
        self.assertEqual({(r[9], r[1]) for r in exported}, {(r["session_id"], r["student_id"]) for r in self.app.record_rows})
        self.assertTrue(all(r[4] == "缺勤" for r in exported))

    def test_semester_dialog_saves_custom_dates_and_can_restore_default(self):
        self.choose(self.app.record_period, self.app.record_period_combo, "当前学期")
        self.assertEqual(len(self.app.record_rows), 9)
        self.app.configure_semester()
        dialog = next(w for w in self.root.winfo_children() if isinstance(w, tk.Toplevel))
        def descendants(widget):
            for child in widget.winfo_children():
                yield child
                yield from descendants(child)
        entries = [w for w in descendants(dialog) if isinstance(w, ttk.Entry)]
        for entry in entries:
            entry.delete(0, "end"); entry.insert(0, "2026-10-21")
        save = next(w for w in descendants(dialog) if isinstance(w, tk.Button) and w.cget("text") == "保存设置")
        save.invoke()
        self.assertEqual(len(self.app.record_rows), 6)
        self.assertEqual(AttendanceStore(self.store.path, seed_path=None).semester_dates(self.now), (date(2026, 10, 21), date(2026, 10, 21)))
        self.app.configure_semester()
        dialog = next(w for w in self.root.winfo_children() if isinstance(w, tk.Toplevel))
        reset = next(w for w in descendants(dialog) if isinstance(w, tk.Button) and w.cget("text") == "恢复默认划分")
        reset.invoke()
        self.assertEqual(len(self.app.record_rows), 9)

    def test_saved_settings_survive_refresh_failure_and_error_uses_live_parent(self):
        self.app.configure_semester()
        dialog = next(w for w in self.root.winfo_children() if isinstance(w, tk.Toplevel))
        def descendants(widget):
            for child in widget.winfo_children():
                yield child
                yield from descendants(child)
        save = next(w for w in descendants(dialog) if isinstance(w, tk.Button) and w.cget("text") == "保存设置")
        with patch.object(self.store, "attendance_records", side_effect=sqlite3.OperationalError("database is locked")), \
                patch("admin_app.messagebox.showerror") as error:
            save.invoke()
        self.assertFalse(dialog.winfo_exists())
        error.assert_called_once()
        self.assertIs(error.call_args.kwargs["parent"], self.root)
        self.app.tick()
        self.assertEqual(len(self.app.record_rows), 6)
        self.assertIn("当前学期日期已保存", self.app.status.get())
        self.assertFalse(self.callback_errors)


if __name__ == "__main__":
    unittest.main()
