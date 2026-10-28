"""四时间边界、旧库迁移、备份兼容与课次界面回归。"""

import json
import sqlite3
import tempfile
import tkinter as tk
import unittest
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from classroom_attendance.ui.admin_app import AdminApp
from classroom_attendance.storage.attendance_store import AttendanceStore
from _face_test_support import authorize_faces
from classroom_attendance.storage.backup_service import SnapshotSource, create_backup, database_schema, database_summary, inspect_backup, restore_backup
from classroom_attendance.courses.course_schedule import course_times, expand_course, normalize_course
from classroom_attendance.ui.student_app import StudentApp


class TimeFixture(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='classroom-session-times-')
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.store = AttendanceStore(self.folder / 'source' / 'attendance.db', seed_path=None)
        for index in range(1, 7):
            self.store.save_student(f'{index:03d}', f'合成学生{index}', '测试班')
        self.now = datetime.now().replace(second=0, microsecond=0)
        authorize_faces(self.store)
        self.times = dict(opens_at=self.now-timedelta(minutes=5), starts_at=self.now+timedelta(minutes=15),
                          late_at=self.now+timedelta(minutes=25), ends_at=self.now+timedelta(minutes=75))
        self.definition = dict(name='四时间模板', course_name='四时间课程', group_name='测试班',
                               opens_time='07:40', start_time='08:00', late_time='08:05', end_time='09:30',
                               next_day=False, opens_previous_day=False, late_next_day=False,
                               late_minutes=10, interval_weeks=1, weekdays=[0], student_ids=None)

    def session(self, **overrides):
        return self.store.create_session('时间测试', '测试班', **dict(self.times, **overrides))

    def downgrade(self):
        with self.store.connect(write=True) as connection:
            connection.execute('ALTER TABLE sessions DROP COLUMN opens_at')
            self.store._set(connection, 'schema_version', '3')

    def backup(self, source=None):
        output = self.folder / 'backup.zip'
        create_backup(source or self.store, output, model=self.folder/'none.yml', manifest=self.folder/'none.json', dataset=self.folder/'no-samples')
        return output


class SessionTimeTests(TimeFixture):
    def test_signin_boundaries_distinguish_open_class_late_and_close(self):
        session_id = self.session()
        with self.assertRaises(ValueError):
            self.store.open_session(session_id, self.times['opens_at']-timedelta(seconds=1))
        self.store.open_session(session_id, self.times['opens_at'])
        with self.assertRaises(ValueError):
            self.store.record_sign_in('006', session_id, self.times['opens_at']-timedelta(seconds=1))
        for student, moment, status in [
            ('001', self.times['opens_at'], 'present'), ('002', self.times['starts_at'], 'present'),
            ('003', self.times['late_at']-timedelta(seconds=1), 'present'), ('004', self.times['late_at'], 'late'),
            ('005', self.times['ends_at']-timedelta(seconds=1), 'late')]:
            self.assertEqual(self.store.record_sign_in(student, session_id, moment)['status'], status)
        with self.assertRaises(ValueError):
            self.store.record_sign_in('006', session_id, self.times['ends_at'])

    def test_manual_open_and_unique_open_session_apply_before_official_class(self):
        first, second = self.session(), self.session()
        self.store.open_manual_session(first, '合成人工开课', now=self.now)
        self.assertEqual(self.store.active_session(self.now)['id'], first)
        self.assertEqual(self.store.teacher_dashboard(self.now)['active']['id'], first)
        with self.assertRaises(ValueError):
            self.store.open_manual_session(second, '另一课次', now=self.now)

    def test_equal_class_and_late_still_accepts_early_present_then_late(self):
        session_id = self.session(late_at=self.times['starts_at'])
        self.store.open_session(session_id, self.times['opens_at'])
        self.assertEqual(self.store.record_sign_in('001', session_id, self.times['opens_at'])['status'], 'present')
        self.assertEqual(self.store.record_sign_in('002', session_id, self.times['starts_at'])['status'], 'late')

    def test_late_at_close_never_marks_accepted_checkin_late(self):
        session_id = self.session(late_at=self.times['ends_at'])
        self.store.open_session(session_id, self.now)
        record = self.store.record_sign_in('001', session_id, self.times['ends_at']-timedelta(seconds=1))
        self.assertEqual(record['status'], 'present')

    def test_invalid_order_timezone_and_non_datetime_fail_without_sessions(self):
        for changed in [dict(opens_at=self.times['starts_at']+timedelta(seconds=1)),
                        dict(late_at=self.times['starts_at']-timedelta(seconds=1)),
                        dict(late_at=self.times['ends_at']+timedelta(seconds=1)),
                        dict(ends_at=self.times['starts_at']), dict(starts_at='2026-10-24 08:00'),
                        dict(starts_at=None), dict(ends_at=None),
                        dict(opens_at=self.now.replace(tzinfo=timezone.utc))]:
            with self.subTest(changed=changed), self.assertRaises(ValueError):
                self.session(**changed)
        self.assertEqual(self.store.sessions(), [])

    def test_maximum_duration_counts_from_open_including_previous_day(self):
        start = self.now
        self.session(opens_at=start-timedelta(hours=1), starts_at=start, late_at=start+timedelta(minutes=1),
                     ends_at=start+timedelta(hours=23))
        with self.assertRaises(ValueError):
            self.session(opens_at=start-timedelta(hours=1, seconds=1), starts_at=start,
                         late_at=start+timedelta(minutes=1), ends_at=start+timedelta(hours=23))

    def test_explicit_late_time_is_not_limited_by_legacy_grace(self):
        session_id = self.session(late_at=self.now+timedelta(hours=4), ends_at=self.now+timedelta(hours=5))
        self.assertEqual(self.store.session(session_id)['late_at'], (self.now+timedelta(hours=4)).isoformat())

    def test_legacy_api_keeps_grace_zero_and_clamped_to_close(self):
        for grace in (0, 10, 180):
            session_id = self.store.create_session('旧接口', '测试班', self.now, self.now+timedelta(minutes=30), grace)
            row = self.store.session(session_id)
            self.assertEqual(row['opens_at'], row['starts_at'])
            self.assertEqual(row['late_at'], (self.now+timedelta(minutes=min(grace, 30))).isoformat())

    def test_formal_date_drives_reports_and_dashboard_for_previous_day_open(self):
        start = self.now.replace(hour=0, minute=10)
        session_id = self.session(opens_at=start-timedelta(minutes=30), starts_at=start,
                                  late_at=start+timedelta(minutes=5), ends_at=start+timedelta(minutes=50))
        self.store.open_session(session_id, start-timedelta(minutes=30))
        self.store.record_sign_in('001', session_id, start-timedelta(minutes=20))
        self.store.close_session(session_id)
        official_day = start.date().isoformat()
        previous_day = (start.date()-timedelta(days=1)).isoformat()
        report = self.store.attendance_report(official_day, official_day, now=start+timedelta(hours=1))
        self.assertEqual(report['totals']['arrived'], 1)
        self.assertEqual(self.store.attendance_report(previous_day, previous_day)['sessions'], [])
        self.assertEqual([row['id'] for row in self.store.teacher_dashboard(start+timedelta(hours=1))['today']], [session_id])

    def test_passing_open_time_does_not_open_draft_or_create_absence(self):
        session_id = self.session()
        self.assertIsNone(self.store.active_session(self.now))
        self.assertEqual(self.store.session(session_id)['state'], 'draft')
        day = self.times['starts_at'].date().isoformat()
        self.assertEqual(self.store.attendance_report(day, day, now=self.times['ends_at'])['totals']['expected'], 0)

    def test_creation_audit_records_all_four_times(self):
        session_id = self.session()
        event = next(row for row in self.store.audit_rows(session_id) if row['action'] == '创建课次')
        self.assertEqual(json.loads(event['after_json']), {key: value.isoformat() for key, value in self.times.items()})

    def test_all_export_forms_include_four_times_and_keep_original_evidence(self):
        session_id = self.session()
        self.store.open_session(session_id, self.now)
        self.store.record_sign_in('001', session_id, self.now)
        self.store.close_session(session_id)
        paths = [self.folder / (name+'.xlsx') for name in ('single', 'records', 'reports')]
        self.store.export_excel(paths[0], session_id)
        begin = self.times['starts_at'].replace(hour=0, minute=0)
        self.store.export_attendance(paths[1], begin, begin+timedelta(days=1))
        day = self.times['starts_at'].date().isoformat()
        self.store.export_attendance_report(paths[2], self.store.attendance_report(day, day))
        for path in paths:
            workbook = load_workbook(path)
            try:
                page = workbook['考勤明细'] if '考勤明细' in workbook.sheetnames else workbook.active
                headers = [cell.value for cell in page[4]]
                self.assertTrue({'开放签到时间', '正式上课时间', '迟到起点', '关闭签到时间'} <= set(headers))
                self.assertEqual(page.cell(5, headers.index('开放签到时间')+1).value,
                                 self.times['opens_at'].strftime('%Y-%m-%d %H:%M:%S') if '考勤明细' in workbook.sheetnames else self.times['opens_at'])
            finally:
                workbook.close()
        self.assertEqual(self.store.session_rows(session_id)[0]['original_signed_at'], self.now.isoformat())


class SessionTimeMigrationTests(TimeFixture):
    def test_v3_upgrade_preserves_times_evidence_password_template_and_review(self):
        session_id = self.store.create_session('旧课次', '测试班', self.now-timedelta(minutes=5), self.now+timedelta(hours=1), 10)
        self.store.open_session(session_id, self.now)
        self.store.record_sign_in('001', session_id, self.now)
        self.store.close_session(session_id)
        self.store.setup_admin('Example-migration-password')
        legacy = {key: value for key, value in self.definition.items() if key not in {'opens_time', 'late_time', 'opens_previous_day', 'late_next_day'}}
        template = self.store.save_course_template(legacy)
        plan = self.store.preview_session_review(session_id)
        self.store.mark_session_reviewed(session_id, plan['token'], note='旧核对')
        before = self.store.session(session_id)
        audit_count = len(self.store.audit_rows())
        self.downgrade()
        migrated = AttendanceStore(self.store.path, seed_path=None)
        self.assertEqual(database_schema(migrated.path), 5)
        self.assertEqual(migrated.session(session_id)['opens_at'], before['starts_at'])
        self.assertEqual(migrated.session(session_id)['late_at'], before['late_at'])
        self.assertEqual(migrated.session_rows(session_id)[0]['original_signed_at'], self.now.isoformat())
        self.assertTrue(migrated.preview_session_review(session_id)['reviewed'])
        self.assertEqual(migrated.course_templates(), [template])
        self.assertEqual(len(migrated.audit_rows()), audit_count)
        self.assertTrue(migrated.authenticate('Example-migration-password'))

    def test_migration_failure_rolls_back_new_column_and_version(self):
        self.store.create_session('旧课次', '测试班', self.now, self.now+timedelta(hours=1))
        self.downgrade()
        with self.store.connect(write=True) as connection:
            connection.execute("CREATE TRIGGER fail_upgrade BEFORE UPDATE ON sessions BEGIN SELECT RAISE(ABORT,'synthetic'); END")
        with self.assertRaises(sqlite3.Error):
            AttendanceStore(self.store.path, seed_path=None)
        with self.store.connect() as connection:
            self.assertNotIn('opens_at', {row[1] for row in connection.execute('PRAGMA table_info(sessions)')})
            self.assertEqual(self.store._get(connection, 'schema_version'), '3')

    def test_schema3_archive_is_read_only_verified_then_upgraded_on_open(self):
        session_id = self.store.create_session('旧课次', '测试班', self.now, self.now+timedelta(hours=1), 10)
        self.downgrade()
        archive = self.backup(SnapshotSource(self.store.path))
        summary = inspect_backup(archive)
        self.assertEqual(summary['schema_version'], 3)
        self.assertEqual(database_schema(self.store.path), 3)
        destination = self.folder / 'restored'
        restore_backup(archive, destination, expected_sha256=summary['archive_sha256'])
        migrated = AttendanceStore(destination/'attendance.db', seed_path=None)
        self.assertEqual(migrated.session(session_id)['opens_at'], self.now.isoformat())
        self.assertEqual(database_schema(migrated.path), 5)

    def test_schema4_backup_and_reopen_keep_independent_open_time(self):
        session_id = self.session()
        reopened = AttendanceStore(self.store.path, seed_path=None)
        self.assertEqual(reopened.session(session_id)['opens_at'], self.times['opens_at'].isoformat())
        archive = self.backup()
        summary = inspect_backup(archive)
        self.assertEqual(summary['schema_version'], 5)
        destination = self.folder / 'restored'
        restore_backup(archive, destination, expected_sha256=summary['archive_sha256'])
        restored = AttendanceStore(destination/'attendance.db', seed_path=None)
        self.assertEqual({key: restored.session(session_id)[key] for key in self.times}, {key: value.isoformat() for key, value in self.times.items()})

    def test_unknown_schema_and_false_schema4_structure_are_rejected(self):
        self.downgrade()
        with self.store.connect(write=True) as connection:
            self.store._set(connection, 'schema_version', '4')
        with self.assertRaises(ValueError):
            database_summary(self.store.path)
        with self.store.connect(write=True) as connection:
            self.store._set(connection, 'schema_version', '999')
        with self.assertRaises(ValueError):
            AttendanceStore(self.store.path, seed_path=None)
        with self.store.connect() as connection:
            self.assertEqual(self.store._get(connection, 'schema_version'), '999')

    def test_backup_rejects_invalid_time_order(self):
        session_id = self.session()
        with self.store.connect(write=True) as connection:
            connection.execute('UPDATE sessions SET opens_at=ends_at WHERE id=?', (session_id,))
        with self.assertRaisesRegex(ValueError, '时间'):
            self.backup()
        self.assertFalse((self.folder/'backup.zip').exists())


class FourTimeScheduleTests(TimeFixture):
    def test_cross_midnight_cycle_uses_explicit_late_on_next_day(self):
        definition = dict(self.definition, opens_time='23:00', start_time='23:30', late_time='00:05', end_time='00:30',
                          late_next_day=True, next_day=True)
        row = expand_course(definition, '2027-03-01', '2027-03-01')[0]
        self.assertEqual([row[key] for key in ('opens_at', 'starts_at', 'late_at', 'ends_at')],
                         ['2027-03-01T23:00:00', '2027-03-01T23:30:00', '2027-03-02T00:05:00', '2027-03-02T00:30:00'])

    def test_previous_day_open_keeps_official_weekday_and_date(self):
        definition = dict(self.definition, opens_time='23:40', opens_previous_day=True, start_time='00:10', late_time='00:20', end_time='01:00')
        row = expand_course(definition, '2027-03-01', '2027-03-01')[0]
        self.assertEqual(row['date'], '2027-03-01')
        self.assertEqual(row['opens_at'], '2027-02-28T23:40:00')

    def test_conflict_in_early_open_window_counts_and_adjacent_is_allowed(self):
        self.store.create_session('占用', '测试班', datetime(2027, 3, 1, 7, 30), datetime(2027, 3, 1, 7, 50))
        plan = self.store.preview_course_schedule(self.definition, '2027-03-01', '2027-03-01')
        self.assertEqual(plan['conflict_count'], 1)
        adjacent = dict(self.definition, opens_time='07:50')
        self.assertEqual(self.store.preview_course_schedule(adjacent, '2027-03-01', '2027-03-01')['conflict_count'], 0)

    def test_legacy_definition_preserves_original_grace_and_new_template_is_v2(self):
        legacy = {key: value for key, value in self.definition.items() if key not in {'opens_time', 'late_time', 'opens_previous_day', 'late_next_day'}}
        legacy.update(start_time='23:30', end_time='00:20', next_day=True, late_minutes=180)
        times = course_times(normalize_course(legacy), date(2027, 3, 1))
        self.assertEqual(times['opens_at'], times['starts_at'])
        self.assertEqual(times['late_at'], times['ends_at'])
        template = self.store.save_course_template(self.definition)
        self.assertEqual(template['version'], 2)
        self.assertEqual(self.store.course_templates(), [template])

    def test_time_edits_invalidate_preview_and_generated_times_match(self):
        plan = self.store.preview_course_schedule(self.definition, '2027-03-01', '2027-03-01')
        changed = dict(self.definition, late_time='08:06')
        with self.assertRaises(ValueError):
            self.store.create_course_schedule(changed, '2027-03-01', '2027-03-01', plan['token'])
        session_id = self.store.create_course_schedule(self.definition, '2027-03-01', '2027-03-01', plan['token'])[0]
        for key in self.times:
            self.assertEqual(self.store.session(session_id)[key], plan['rows'][0][key])

    def test_invalid_explicit_times_flags_and_date_overflow_fail(self):
        for values in [dict(opens_time='08:01'), dict(late_time='07:59'), dict(opens_time='bad'),
                       dict(late_time=None), dict(opens_previous_day='yes'), dict(late_next_day=True)]:
            with self.subTest(values=values), self.assertRaises(ValueError):
                normalize_course(dict(self.definition, **values))
        with self.assertRaises(ValueError):
            expand_course(dict(self.definition, opens_previous_day=True, weekdays=[0]), '0001-01-01', '0001-01-01')


class SessionTimeGuiTests(TimeFixture):
    def setUp(self):
        super().setUp()
        self.root = tk.Tk()
        self.errors = patch('classroom_attendance.ui.admin_app.messagebox.showerror').start()
        self.addCleanup(patch.stopall)
        with patch('classroom_attendance.ui.admin_app.model_health', return_value=(False, '未训练', None)):
            self.admin = AdminApp(self.root, self.store, authenticated=True, async_io=False)
        self.addCleanup(self.cleanup)
        self.root.update()

    def cleanup(self):
        self.admin.close()
        try:
            self.root.update_idletasks()
            self.root.destroy()
        except tk.TclError:
            pass

    def test_single_form_saves_four_explicit_times_and_rejects_bad_order(self):
        self.admin.course.set('界面时间课程')
        for variable, key in ((self.admin.opens, 'opens_at'), (self.admin.starts, 'starts_at'),
                              (self.admin.late, 'late_at'), (self.admin.ends, 'ends_at')):
            variable.set(self.times[key].strftime('%Y-%m-%d %H:%M'))
        self.admin.create_session()
        row = self.store.session(self.admin.selected_session_id())
        self.assertEqual({key: row[key] for key in self.times}, {key: value.isoformat() for key, value in self.times.items()})
        self.admin.late.set(self.times['opens_at'].strftime('%Y-%m-%d %H:%M'))
        self.admin.create_session()
        self.assertEqual(len(self.store.sessions()), 1)
        self.assertTrue(self.errors.called)

    def test_student_sees_four_times_before_formal_class(self):
        session_id = self.session()
        self.store.open_session(session_id, self.now)
        root = tk.Toplevel(self.root)
        with patch('classroom_attendance.ui.student_app.model_health', return_value=(True, '就绪', {})):
            student = StudentApp(root, self.store, async_io=False)
        self.root.update_idletasks()
        self.assertEqual(student.current_session['id'], session_id)
        for text in ('开放签到', '正式上课', '迟到起点', '关闭签到'):
            self.assertIn(text, student.details.cget('text'))
        student.close()

    def test_cycle_preview_shows_four_times_and_legacy_template_loads_exactly(self):
        legacy = dict(name='旧模板', course_name='旧课', group_name='测试班', start_time='23:30', end_time='00:20',
                      next_day=True, late_minutes=180, interval_weeks=1, weekdays=[0], student_ids=None)
        self.store.save_course_template(legacy)
        self.admin.course_schedule()
        dialog = self.admin.schedule_dialog
        dialog.variables['template_choice'].set(next(iter(dialog.template_choices)))
        dialog.load_selected()
        self.assertEqual(dialog.variables['opens_time'].get(), '23:30')
        self.assertEqual(dialog.variables['late_time'].get(), '00:20')
        self.assertTrue(dialog.late_next_day.get())
        dialog.variables['first_date'].set('2027-03-01')
        dialog.variables['last_date'].set('2027-03-01')
        dialog.preview()
        self.assertEqual(len(dialog.preview_tree['columns']), 7)
        self.assertEqual(dialog.plan['rows'][0]['late_at'], '2027-03-02T00:20:00')

    def test_small_form_keeps_create_fixed_and_all_time_fields_reachable(self):
        self.admin.notebook.select(self.admin.pages['sessions'])
        self.root.geometry('940x600')
        self.root.update()
        self.admin.session_form_scroll.canvas.yview_moveto(1)
        self.root.update()
        control = self.admin.create_session_button
        self.assertTrue(control.winfo_ismapped())
        self.assertLessEqual(control.winfo_rooty()+control.winfo_height(), self.root.winfo_rooty()+self.root.winfo_height())
        self.assertGreater(self.admin.sessions_tree.winfo_height(), 100)

    def test_historical_navigation_uses_formal_date_not_previous_open_day(self):
        start = self.now.replace(hour=0, minute=10)
        session_id = self.session(opens_at=start-timedelta(minutes=30), starts_at=start,
                                  late_at=start+timedelta(minutes=5), ends_at=start+timedelta(hours=1))
        self.store.open_session(session_id, start-timedelta(minutes=30))
        self.store.close_session(session_id)
        self.admin.show_dashboard_records(session_id)
        self.assertEqual(self.admin.record_target['starts_at'][:10], start.date().isoformat())
        self.assertEqual({row['session_id'] for row in self.admin.record_rows}, {session_id})


if __name__ == '__main__':
    unittest.main()
