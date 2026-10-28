"""统计口径、历史快照、一致导出与教师报表页回归。"""

import sqlite3
import tempfile
import threading
import time
import tkinter as tk
import unittest
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from classroom_attendance.ui.admin_app import AdminApp
from classroom_attendance.core.app_errors import TaskCancelled
from classroom_attendance.storage.attendance_store import AttendanceStore
from _face_test_support import authorize_faces
from classroom_attendance.attendance.attendance_reports import rate_text
from classroom_attendance.core.task_runtime import TaskControl, current_control


class ReportFixture(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='attendance-reports-')
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.store = AttendanceStore(self.folder / 'attendance.db', seed_path=None)
        for student_id, name, group in [('001', '合成甲', '甲班'), ('002', '合成乙', '甲班'), ('003', '合成丙', '乙班')]:
            self.store.save_student(student_id, name, group)
        self.now = datetime(2026, 10, 24, 12)
        authorize_faces(self.store)

    def session(self, day=1, *, course='合成课程', group='全部学生', students=None, state='closed', start=None):
        start = start or datetime(2026, 10, day, 9)
        session_id = self.store.create_session(course, group, start, start + timedelta(hours=1), student_ids=students)
        if state in {'closed', 'open'}:
            self.store.open_session(session_id, start)
        if state == 'closed':
            self.store.close_session(session_id)
        elif state == 'cancelled':
            self.store.cancel_session(session_id, '合成停课')
        return session_id

    def report(self, **kwargs):
        return self.store.attendance_report('2026-10-01', '2026-10-31', now=self.now, **kwargs)

    def mark(self, session_id, student_id, status, **kwargs):
        self.store.correct_attendance(session_id, student_id, status, '合成更正', **kwargs)


class AttendanceReportTests(ReportFixture):
    def test_rates_count_late_and_keep_leave_in_denominator(self):
        first = self.session()
        second = self.session(2)
        self.mark(first, '001', 'present', time_unconfirmed=True)
        self.mark(first, '002', 'late', time_unconfirmed=True)
        self.mark(first, '003', 'leave')
        self.mark(second, '001', 'present', time_unconfirmed=True)
        report = self.report()
        totals = report['totals']
        self.assertEqual([totals[key] for key in ('expected', 'present', 'late', 'leave', 'absent', 'arrived')], [6, 2, 1, 1, 2, 3])
        self.assertEqual(totals['attendance_rate'], .5)
        self.assertEqual(totals['on_time_rate'], 2 / 6)
        self.assertEqual(totals['time_unconfirmed'], 3)
        self.assertEqual(report['students'][1]['attendance_rate'], .5)
        self.assertEqual(rate_text(1 / 32), '3.13%')
        self.assertEqual(rate_text(57 / 800), '7.13%')

    def test_only_ended_sessions_and_exact_expiry_boundary_count_without_writes(self):
        ended = self.session(2)
        self.session(3, state='draft')
        self.session(4, state='cancelled')
        expired = self.session(state='open', start=self.now - timedelta(hours=1))
        before = len(self.store.audit_rows())
        report = self.report()
        self.assertEqual({row['id'] for row in report['sessions']}, {ended, expired})
        self.assertEqual(report['totals']['absent'], 6)
        with self.store.connect() as connection:
            self.assertEqual(connection.execute('SELECT state FROM sessions WHERE id=?', (expired,)).fetchone()[0], 'open')
        self.assertEqual(len(self.store.audit_rows()), before)
        report = self.store.attendance_report('2026-10-01', '2026-10-31', now=self.now-timedelta(seconds=1))
        self.assertEqual([row['id'] for row in report['sessions']], [ended])

    def test_in_progress_signed_and_manual_results_do_not_enter_denominator(self):
        session_id = self.session(state='open', start=self.now - timedelta(minutes=5))
        class ReportClock(datetime):
            @classmethod
            def now(cls, tz=None):
                return datetime(2026, 10, 24, 12, tzinfo=tz)
        with patch('classroom_attendance.storage.attendance_store.datetime', ReportClock):
            self.mark(session_id, '001', 'present', time_unconfirmed=True)
            self.mark(session_id, '002', 'leave')
        self.assertEqual(self.report()['totals']['expected'], 0)

    def test_date_range_includes_last_day_and_assigns_cross_day_by_start(self):
        first = self.session(start=datetime(2026, 9, 30, 23, 30))
        last = self.session(start=datetime(2026, 10, 31, 23, 30))
        self.session(start=datetime(2026, 11, 1))
        result = self.store.attendance_report('2026-10-31', '2026-10-31', now=datetime(2026, 11, 2))
        self.assertEqual([row['id'] for row in result['sessions']], [last])
        self.assertNotIn(first, [row['id'] for row in result['sessions']])

    def test_same_day_different_sessions_and_course_classes_are_distinct(self):
        first = self.session(students=['001'])
        second = self.session(group='甲班')
        report = self.report()
        self.assertEqual(report['totals']['expected'], 3)
        self.assertEqual(len(report['courses']), 2)
        self.assertEqual(set(report['students'][0]['cells']), {first, second})
        self.assertEqual(report['students'][1]['cells'], {second: 'absent'})

    def test_weighted_total_uses_person_sessions_instead_of_average_rates(self):
        first = self.session(students=['001'])
        self.mark(first, '001', 'present', time_unconfirmed=True)
        self.session(2, course='另一课程')
        report = self.report()
        self.assertEqual(report['totals']['attendance_rate'], .25)
        self.assertEqual(sum(row['attendance_rate'] for row in report['courses']) / 2, .5)

    def test_historical_profiles_and_disabled_students_use_snapshots(self):
        first = self.session()
        self.store.save_student('001', '合成新名', '丙班')
        self.store.set_student_active('002', False)
        second = self.session(2)
        report = self.report()
        self.assertEqual(report['totals']['expected'], 5)
        self.assertEqual(report['students'][0]['name'], '合成新名')
        self.assertEqual(report['students'][0]['group_name'], '丙班')
        old = next(row for row in report['rows'] if row['student_id'] == '001' and row['session_id'] == first)
        self.assertEqual((old['name'], old['group_name']), ('合成甲', '甲班'))
        self.assertEqual(report['students'][1]['cells'], {first: 'absent'})
        self.assertIn(second, report['students'][0]['cells'])

    def test_group_and_course_filters_use_exact_historical_student_class(self):
        self.session(course="课程' OR 1=1 --")
        self.store.save_student('001', '合成甲', '丙班')
        self.session(2, course='其他课程')
        result = self.report(course_name="课程' OR 1=1 --", group_name='甲班')
        self.assertEqual({row['student_id'] for row in result['students']}, {'001', '002'})
        self.assertEqual(result['totals']['expected'], 2)
        self.assertEqual(len(result['sessions']), 1)
        self.assertEqual(len(result['options']['courses']), 2)
        self.assertEqual(self.report(course_name='不存在课程')['totals']['expected'], 0)

    def test_roster_removal_preserves_evidence_and_readdition_recounts(self):
        session_id = self.session()
        plan = self.store.preview_roster_change(session_id, remove_ids=['001'])
        self.store.apply_roster_change(session_id, plan['token'], '合成移出', remove_ids=['001'])
        self.assertEqual(self.report()['totals']['expected'], 2)
        plan = self.store.preview_roster_change(session_id, add_ids=['001'])
        self.store.apply_roster_change(session_id, plan['token'], '合成增补', add_ids=['001'])
        self.assertEqual(self.report()['totals']['expected'], 3)

    def test_corrections_use_current_status_keep_original_evidence_and_refresh(self):
        start = datetime(2026, 10, 1, 9)
        session_id = self.session(state='open', start=start)
        self.store.record_sign_in('001', session_id, start + timedelta(minutes=1))
        self.store.close_session(session_id)
        old = self.report()
        self.mark(session_id, '001', 'leave')
        result = self.report()
        self.assertEqual(old['totals']['arrived'], 1)
        self.assertEqual(result['totals']['arrived'], 0)
        self.assertEqual(result['rows'][0]['original_signed_at'], '2026-10-01T09:01:00')
        self.assertEqual(result['rows'][0]['status'], 'leave')

    def test_empty_roster_session_and_empty_scope_have_undefined_rate(self):
        empty = self.report()
        self.assertIsNone(empty['totals']['attendance_rate'])
        session_id = self.session()
        plan = self.store.preview_roster_change(session_id, remove_ids=['001', '002', '003'])
        self.store.apply_roster_change(session_id, plan['token'], '合成空名单', remove_ids=['001', '002', '003'])
        report = self.report()
        self.assertEqual(len(report['sessions']), 1)
        self.assertEqual(report['students'], [])
        self.assertIsNone(report['courses'][0]['attendance_rate'])
        self.assertEqual(self.report(group_name='甲班')['sessions'], [])

    def test_invalid_dates_and_filters_fail_without_broadening(self):
        for first, last in [('2026-02-30', '2026-10-31'), ('2026-10-31', '2026-10-01'), ('2026-10-01', '9999-12-31')]:
            with self.assertRaises(ValueError):
                self.store.attendance_report(first, last)
        for value in ['', ' ', 42]:
            with self.assertRaises(ValueError):
                self.report(course_name=value)

    def test_sessions_and_details_share_one_sqlite_read_snapshot(self):
        self.session(students=['001'])
        original = self.store.connect
        store = self.store

        class Proxy:
            def execute(self, query, *args):
                result = connection.execute(query, *args)
                if query.startswith('SELECT s.*'):
                    result.fetchall()
                    self.test.session(2, students=['002'])
                    self.test.assertTrue(connection.in_transaction)
                    return connection.execute(query, *args)
                return result

        @contextmanager
        def connected():
            with original() as active:
                nonlocal connection
                connection = active
                proxy = Proxy()
                proxy.test = self
                with patch.object(store, 'connect', original):
                    yield proxy

        connection = None
        with patch.object(self.store, 'connect', connected):
            self.assertEqual(self.report()['totals']['expected'], 1)
        self.assertEqual(self.report()['totals']['expected'], 2)


class AttendanceReportExportTests(ReportFixture):
    def test_export_has_five_sheets_numeric_rates_safe_strings_and_matrix_gaps(self):
        self.store.save_student('001', '=HYPERLINK("x")', '=甲班')
        first = self.session(course='=合成课程', students=['001'])
        second = self.session(2, course='另一课程', students=['002'])
        self.mark(first, '001', 'late', time_unconfirmed=True, note='=1+1')
        report = self.report()
        output = self.folder / '报表.xlsx'
        self.store.export_attendance_report(output, report)
        with self.book(output) as workbook:
            self.assertEqual(workbook.sheetnames, ['学生学期汇总', '课程汇总', '考勤矩阵', '考勤明细', '统计口径'])
            students = workbook['学生学期汇总']
            self.assertEqual(students['A5'].value, '001')
            self.assertEqual(students['A5'].data_type, 's')
            self.assertEqual(students['K5'].value, 1)
            self.assertEqual(students['K5'].number_format, '0.00%')
            self.assertEqual(students['K7'].value, .5)
            matrix = workbook['考勤矩阵']
            self.assertEqual((matrix['D5'].value, matrix['E5'].value), ('迟到', '—'))
            self.assertEqual((matrix['D6'].value, matrix['E6'].value), ('—', '缺勤'))
            self.assertIn(f'#{first}', matrix['D4'].value)
            self.assertIn(f'#{second}', matrix['E4'].value)
            for page in workbook:
                for cells in page:
                    self.assertFalse(any(cell.data_type == 'f' for cell in cells))

    @contextmanager
    def book(self, path):
        workbook = load_workbook(path)
        try:
            yield workbook
        finally:
            workbook.close()

    def test_export_uses_supplied_snapshot_after_database_correction(self):
        session_id = self.session()
        snapshot = self.report(group_name='甲班')
        self.mark(session_id, '001', 'leave')
        output = self.folder / '快照.xlsx'
        self.store.export_attendance_report(output, snapshot)
        with self.book(output) as workbook:
            self.assertEqual(workbook['考勤矩阵']['D5'].value, '缺勤')
            self.assertIn('甲班', workbook['考勤矩阵']['A2'].value)
            self.assertEqual(workbook['学生学期汇总']['D7'].value, 2)

    def test_empty_export_explains_zero_denominator(self):
        output = self.folder / '空.xlsx'
        self.store.export_attendance_report(output, self.report())
        with self.book(output) as workbook:
            self.assertEqual(workbook['学生学期汇总']['K5'].value, '—')
            self.assertEqual(workbook['考勤矩阵'].max_row, 4)
            self.assertIn('请假保留在分母', workbook['统计口径']['A6'].value)

    def test_cancel_after_save_preserves_destination_and_removes_temporary(self):
        from openpyxl.workbook.workbook import Workbook
        output = self.folder / '现有.xlsx'
        output.write_bytes(b'original')
        control = TaskControl()
        save = Workbook.save

        def cancel(workbook, filename):
            save(workbook, filename)
            control.cancelled.set()

        token = current_control.set(control)
        try:
            with patch.object(Workbook, 'save', cancel), self.assertRaises(TaskCancelled):
                self.store.export_attendance_report(output, self.report())
        finally:
            current_control.reset(token)
        self.assertEqual(output.read_bytes(), b'original')
        self.assertEqual(list(self.folder.glob('.*.tmp.xlsx')), [])

    def test_failed_save_preserves_destination_and_removes_temporary(self):
        output = self.folder / '现有.xlsx'
        output.write_bytes(b'original')
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=OSError('synthetic')), self.assertRaises(OSError):
            self.store.export_attendance_report(output, self.report())
        self.assertEqual(output.read_bytes(), b'original')
        self.assertEqual(list(self.folder.glob('.*.tmp.xlsx')), [])

    def test_excel_extension_and_matrix_column_limit_are_checked(self):
        report = self.report()
        with self.assertRaises(ValueError):
            self.store.export_attendance_report(self.folder / 'wrong.csv', report)
        report['sessions'] = [None] * 16382
        with self.assertRaisesRegex(ValueError, '上限'):
            self.store.export_attendance_report(self.folder / 'too-wide.xlsx', report)


class AttendanceReportGuiTests(ReportFixture):
    def setUp(self):
        super().setUp()
        self.session()
        self.root = tk.Tk()
        self.errors = []
        self.root.report_callback_exception = lambda *args: self.errors.append(args)
        self.store.set_semester_dates('2026-10-01', '2026-10-31')
        with patch('classroom_attendance.ui.admin_app.model_health', return_value=(False, '未训练', None)):
            self.admin = AdminApp(self.root, self.store, authenticated=True, async_io=False)
        self.panel = self.admin.reports
        self.admin.notebook.select(self.admin.pages['reports'])
        self.root.update()

    def tearDown(self):
        self.admin.close()
        self.assertFalse(self.errors)

    def test_tab_loads_configured_semester_and_all_three_views(self):
        self.assertEqual(self.panel.starts_on.get(), '2026-10-01')
        self.assertEqual(len(self.panel.student_tree.get_children()), 3)
        self.assertEqual(len(self.panel.course_tree.get_children()), 1)
        self.assertEqual(len(self.panel.matrix_tree.get_children()), 3)
        self.assertIn('出勤率 0.00%', self.panel.stats.get())
        self.assertEqual(str(self.panel.export_button['state']), 'normal')

    def test_changed_filters_disable_export_until_regenerated(self):
        self.panel.group.set('班级：甲班')
        self.assertEqual(str(self.panel.export_button['state']), 'disabled')
        self.panel.refresh()
        self.assertEqual(self.panel.snapshot['totals']['expected'], 2)
        self.assertEqual(str(self.panel.export_button['state']), 'normal')
        output = self.folder / 'gui.xlsx'
        with patch('classroom_attendance.attendance.attendance_reports_ui.filedialog.asksaveasfilename', return_value=str(output)), patch.object(self.admin, 'run_job') as run:
            self.panel.export()
        run.call_args.args[1]()
        self.assertTrue(output.is_file())
        self.assertEqual(self.panel.snapshot['group_name'], '甲班')

    def test_failure_disables_old_snapshot_then_retry_recovers(self):
        with patch.object(self.store, 'attendance_report', side_effect=sqlite3.OperationalError('busy')):
            self.panel.refresh()
        self.assertIsNone(self.panel.snapshot)
        self.assertEqual(str(self.panel.export_button['state']), 'disabled')
        self.panel.refresh()
        self.assertIsNotNone(self.panel.snapshot)
        self.assertEqual(str(self.panel.export_button['state']), 'normal')

    def test_invalid_dates_never_export_old_display(self):
        self.panel.starts_on.set('错误日期')
        self.panel.refresh()
        self.assertIsNone(self.panel.snapshot)
        with patch('classroom_attendance.attendance.attendance_reports_ui.filedialog.asksaveasfilename') as choose:
            self.panel.export()
        choose.assert_not_called()
        self.panel.use_semester()
        self.assertEqual(self.panel.snapshot['starts_on'], '2026-10-01')
        self.panel.starts_on.set('20261001')
        self.panel.refresh()
        self.assertEqual(self.panel.starts_on.get(), '2026-10-01')
        self.assertEqual(str(self.panel.export_button['state']), 'normal')

    def test_matrix_student_detail_uses_snapshot(self):
        self.panel.matrix_tree.selection_set('001')
        with patch.object(self.admin, 'text_details') as show:
            self.panel.show_student()
            self.assertIn('缺勤', show.call_args.args[1])
            self.assertIn('001', show.call_args.args[1])
            self.panel.show_rules()
            self.assertIn('请假保留在分母', show.call_args.args[1])

    def test_small_screen_keeps_actions_table_and_horizontal_scroll(self):
        self.root.geometry('940x600')
        self.root.update()
        for widget in (self.panel.generate_button, self.panel.export_button, self.panel.semester_button):
            self.assertTrue(widget.winfo_ismapped())
            self.assertLessEqual(widget.winfo_rootx() + widget.winfo_width(), self.root.winfo_rootx() + self.root.winfo_width())
            self.assertLessEqual(widget.winfo_rooty() + widget.winfo_height(), self.root.winfo_rooty() + self.root.winfo_height())
        self.assertGreater(self.panel.student_tree.winfo_height(), 100)
        self.panel.student_tree.xview_moveto(1)
        self.root.update()
        self.assertAlmostEqual(self.panel.student_tree.xview()[1], 1, places=2)

    def test_slow_query_keeps_ui_responsive_and_discards_changed_filters(self):
        self.panel.async_io = True
        original = self.store.attendance_report
        started, release = threading.Event(), threading.Event()

        def slow(*args, **kwargs):
            started.set()
            release.wait(2)
            return original(*args, **kwargs)

        with patch.object(self.store, 'attendance_report', side_effect=slow):
            self.panel.refresh()
            self.assertTrue(started.wait(1))
            heartbeat = []
            self.root.after(10, lambda: heartbeat.append(True))
            self.panel.group.set('班级：乙班')
            release.set()
            self.pump()
        self.assertTrue(heartbeat)
        self.assertEqual(self.panel.snapshot['group_name'], None)
        self.assertEqual(str(self.panel.export_button['state']), 'disabled')

    def pump(self):
        deadline = time.monotonic() + 3
        while self.panel.loading and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.005)
        self.assertFalse(self.panel.loading)


if __name__ == '__main__':
    unittest.main()
