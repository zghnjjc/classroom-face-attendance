"""本地考勤数据：学生、课次名单快照、签到、管理员验证与操作记录。"""

import csv
import hashlib
import hmac
import json
import secrets
import sqlite3
import tempfile
import zipfile
from contextlib import closing, contextmanager
from datetime import datetime, timedelta
from math import ceil
from pathlib import Path

from app_config import DATABASE_PATH, SEED_PATH
from attendance_fields import (STATUS_TEXT, ARRIVED_STATUSES, ARRIVED_STATUS_SQL,
                               is_time_unknown, arrival_time_note, excel_time, attendance_base_values)
from audit_actions import (CORRECT_ATTENDANCE, FACE_CHECKIN, SAVE_FACE_EVIDENCE,
                           CHANGE_ROSTER, CLOSE_SESSION, EXPIRE_SESSION, OPEN_SESSION, CANCEL_SESSION)
from attendance_filters import default_semester_dates, validate_semester_dates
from app_config import DB_READ_TIMEOUT, DB_WRITE_TIMEOUT, PASSWORD_ITERATIONS, LEGACY_PASSWORD_ITERATIONS
from app_errors import ValidationError
from task_runtime import checkpoint
from course_schedule import CourseScheduleStore
from teacher_dashboard import TeacherDashboardStore
from attendance_reports import AttendanceReportStore
from session_times import SCHEMA_VERSION, resolve_session_times
from maintenance_tools import require_disk_space
from face_lifecycle import initialize_schema, require_authorized

UNSET = object()

# 明细、汇总与导出共同使用的状态判定；只有真正开启过的课次才能推导缺勤。
ATTENDANCE_ROW_SQL = """SELECT ss.session_id,ss.student_id,ss.name,ss.group_name,ss.included,
    s.course_name,s.group_name AS session_group_name,s.starts_at AS session_starts_at,
    s.opens_at AS session_opens_at,s.late_at AS session_late_at,s.ends_at AS session_ends_at,s.state AS session_state,
    CASE WHEN ss.included=0 THEN 'excluded'
         WHEN s.state='draft' THEN 'not_started'
         WHEN s.state='cancelled' THEN 'cancelled'
         ELSE COALESCE(c.status,CASE WHEN s.state='closed' OR (s.state='open' AND s.ends_at<=:now)
              THEN 'absent' ELSE 'pending' END) END AS status,
    c.signed_at,c.method,COALESCE(c.note,'') AS note,
    e.signed_at AS original_signed_at,e.method AS original_method,e.origin AS evidence_origin
    FROM sessions s JOIN session_students ss ON ss.session_id=s.id
    LEFT JOIN checkins c ON c.session_id=ss.session_id AND c.student_id=ss.student_id
    LEFT JOIN face_evidence e ON e.session_id=ss.session_id AND e.student_id=ss.student_id"""


def stamp(value=None):
    return (value or datetime.now()).replace(microsecond=0).isoformat()


class AttendanceStore(CourseScheduleStore, TeacherDashboardStore, AttendanceReportStore):
    attendance_row_sql = ATTENDANCE_ROW_SQL
    def __init__(self, path=DATABASE_PATH, seed_path=SEED_PATH):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        # 已有数据库先只读检查；拒绝未知版本时不能先改变 journal_mode 或结构。
        if self.path.is_file() and self.path.stat().st_size:
            with closing(sqlite3.connect(self.path.resolve().as_uri() + '?mode=ro', uri=True,
                                        timeout=DB_READ_TIMEOUT)) as existing:
                if existing.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='metadata'").fetchone():
                    row = existing.execute("SELECT value FROM metadata WHERE key='schema_version'").fetchone()
                    try:
                        version = int(row[0]) if row else 0
                    except (TypeError, ValueError):
                        raise ValueError('数据库版本无效，请使用对应版本的软件，未迁移数据。') from None
                    if not 0 <= version <= SCHEMA_VERSION:
                        raise ValueError('此数据库版本高于当前程序或无效，请使用对应版本的软件，未迁移数据。')
        with self.connect() as connection:
            connection.execute("PRAGMA journal_mode=WAL")
            connection.executescript("""
                CREATE TABLE IF NOT EXISTS students(student_id TEXT PRIMARY KEY, name TEXT NOT NULL,
                    group_name TEXT NOT NULL DEFAULT '未分班', active INTEGER NOT NULL DEFAULT 1);
                CREATE TABLE IF NOT EXISTS metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS attendance(id INTEGER PRIMARY KEY, student_id TEXT NOT NULL,
                    day TEXT NOT NULL, signed_at TEXT NOT NULL, UNIQUE(student_id,day));
                CREATE TABLE IF NOT EXISTS sessions(id INTEGER PRIMARY KEY, course_name TEXT NOT NULL,
                    group_name TEXT NOT NULL, starts_at TEXT NOT NULL, late_at TEXT NOT NULL,
                    ends_at TEXT NOT NULL, state TEXT NOT NULL DEFAULT 'draft', created_at TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS session_students(session_id INTEGER NOT NULL REFERENCES sessions(id),
                    student_id TEXT NOT NULL REFERENCES students(student_id), name TEXT NOT NULL,
                    group_name TEXT NOT NULL, PRIMARY KEY(session_id,student_id));
                CREATE TABLE IF NOT EXISTS checkins(session_id INTEGER NOT NULL, student_id TEXT NOT NULL,
                    status TEXT NOT NULL, signed_at TEXT, method TEXT NOT NULL, note TEXT NOT NULL DEFAULT '',
                    PRIMARY KEY(session_id,student_id),
                    FOREIGN KEY(session_id,student_id) REFERENCES session_students(session_id,student_id));
                CREATE TABLE IF NOT EXISTS audit_log(id INTEGER PRIMARY KEY, created_at TEXT NOT NULL,
                    action TEXT NOT NULL, session_id INTEGER, student_id TEXT, detail TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS idx_checkins_session ON checkins(session_id);
                CREATE INDEX IF NOT EXISTS idx_sessions_starts_at ON sessions(starts_at);
                CREATE TABLE IF NOT EXISTS face_evidence(session_id INTEGER NOT NULL, student_id TEXT NOT NULL,
                    signed_at TEXT NOT NULL, method TEXT NOT NULL, recorded_at TEXT NOT NULL, origin TEXT NOT NULL,
                    PRIMARY KEY(session_id,student_id),
                    FOREIGN KEY(session_id,student_id) REFERENCES session_students(session_id,student_id));
                CREATE TRIGGER IF NOT EXISTS preserve_face_evidence_update BEFORE UPDATE ON face_evidence
                    BEGIN SELECT RAISE(ABORT,'原始人脸签到证据不可修改'); END;
                CREATE TRIGGER IF NOT EXISTS preserve_face_evidence_delete BEFORE DELETE ON face_evidence
                    BEGIN SELECT RAISE(ABORT,'原始人脸签到证据不可删除'); END;
            """)
            connection.execute("BEGIN IMMEDIATE")
            version = int(self._get(connection, "schema_version", "0"))
            if not 0 <= version <= SCHEMA_VERSION:
                raise ValueError('此数据库版本高于当前程序或无效，请使用对应版本的软件，未迁移数据。')
            columns = {row[1] for row in connection.execute("PRAGMA table_info(students)")}
            if "group_name" not in columns:
                connection.execute("ALTER TABLE students ADD COLUMN group_name TEXT NOT NULL DEFAULT '未分班'")
            if "active" not in columns:
                connection.execute("ALTER TABLE students ADD COLUMN active INTEGER NOT NULL DEFAULT 1")
            additions = {
                "sessions": {"cancel_reason": "TEXT NOT NULL DEFAULT ''", "cancelled_at": "TEXT",
                             "opens_at": "TEXT NOT NULL DEFAULT ''"},
                "session_students": {"included": "INTEGER NOT NULL DEFAULT 1"},
                "audit_log": {"actor_label": "TEXT NOT NULL DEFAULT ''", "actor_kind": "TEXT NOT NULL DEFAULT 'legacy_unknown'",
                              "before_json": "TEXT", "after_json": "TEXT"},
            }
            for table, fields in additions.items():
                existing = {row[1] for row in connection.execute(f"PRAGMA table_info({table})")}
                for field, declaration in fields.items():
                    if field not in existing:
                        connection.execute(f"ALTER TABLE {table} ADD COLUMN {field} {declaration}")
            if not self._get(connection, "seed_imported") and seed_path and Path(seed_path).exists():
                data = json.loads(Path(seed_path).read_text(encoding="utf-8"))
                for student in data.get("students", []):
                    connection.execute("INSERT OR IGNORE INTO students(student_id,name,group_name) VALUES(?,?,?)",
                                       (student["student_id"], student["name"], student.get("group_name", "未分班")))
                for record in data.get("attendance", []):
                    connection.execute("INSERT OR IGNORE INTO attendance(student_id,day,signed_at) VALUES(?,?,?)",
                                       (record["student_id"], record["signed_at"][:10], record["signed_at"]))
                self._set(connection, "seed_imported", "1")
            if version < 2:
                self._migrate_daily_records(connection)
            if version < 3:
                # 只能保存升级时仍存在的原值，旧版本已覆盖的刷脸时间无法反推。
                connection.execute("""INSERT OR IGNORE INTO face_evidence
                    SELECT session_id,student_id,signed_at,'face',?,'migrated_current'
                    FROM checkins WHERE method='face' AND signed_at IS NOT NULL""", (stamp(),))
                self._set(connection, "schema_version", "3")
            if version < 4:
                connection.execute('UPDATE sessions SET opens_at=starts_at')
                self._set(connection, 'schema_version', '4')
            if version < 5:
                initialize_schema(connection)
                self._set(connection, 'schema_version', str(SCHEMA_VERSION))
        from face_delete_recovery import recover_face_deletions
        recover_face_deletions(self)

    @contextmanager
    def connect(self, write=False):
        timeout = DB_WRITE_TIMEOUT if write else DB_READ_TIMEOUT
        connection = sqlite3.connect(self.path, timeout=timeout)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys=ON")
        connection.execute(f"PRAGMA busy_timeout={int(timeout*1000)}")
        try:
            with connection:
                yield connection
                checkpoint()  # 取消或超时后绝不提交迟到的写入。
        finally:
            connection.close()

    @staticmethod
    def _get(connection, key, default=None):
        row = connection.execute("SELECT value FROM metadata WHERE key=?", (key,)).fetchone()
        return row[0] if row else default

    @staticmethod
    def _set(connection, key, value):
        connection.execute("INSERT INTO metadata VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))

    @staticmethod
    def _audit(connection, action, session_id=None, student_id=None, detail="", *, actor="", actor_kind=None,
               before=None, after=None):
        actor = actor.strip()
        if len(actor) > 80:
            raise ValueError("操作署名不超过 80 字。")
        kind = actor_kind or ("management_signature" if actor else "management_entry")
        encode = lambda value: json.dumps(value, ensure_ascii=False, sort_keys=True) if value is not None else None
        connection.execute("""INSERT INTO audit_log
            (created_at,action,session_id,student_id,detail,actor_label,actor_kind,before_json,after_json)
            VALUES(?,?,?,?,?,?,?,?,?)""",
            (stamp(), action, session_id, student_id, detail, actor, kind, encode(before), encode(after)))

    @staticmethod
    def _session(connection, session_id):
        row = connection.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
        if not row:
            raise ValueError("课次不存在。")
        return dict(row)

    @staticmethod
    def _reason(reason):
        reason = reason.strip()
        if not reason or len(reason) > 300:
            raise ValueError("请填写原因（1 至 300 字）。")
        return reason

    def _migrate_daily_records(self, connection):
        for row in connection.execute("SELECT DISTINCT day FROM attendance ORDER BY day").fetchall():
            day = row[0]
            session_id = connection.execute(
                "INSERT INTO sessions(course_name,group_name,starts_at,late_at,ends_at,state,created_at) VALUES(?,?,?,?,?,'closed',?)",
                ("历史签到", "原项目名单", f"{day}T00:00:00", f"{day}T23:59:59", f"{day}T23:59:59", stamp()),
            ).lastrowid
            connection.execute("INSERT INTO session_students(session_id,student_id,name,group_name) SELECT ?,student_id,name,group_name FROM students", (session_id,))
            connection.execute("""INSERT INTO checkins(session_id,student_id,status,signed_at,method,note)
                SELECT ?,student_id,'present',signed_at,'legacy','从原项目迁移' FROM attendance WHERE day=?""", (session_id, day))
            self._audit(connection, "迁移历史记录", session_id, detail=day)

    @staticmethod
    def validate_student(student_id, name, group_name="未分班"):
        student_id, name, group_name = str(student_id).strip(), str(name).strip(), str(group_name).strip() or "未分班"
        if not student_id.isascii() or not student_id.isdigit() or len(student_id) > 24:
            raise ValueError("学号必须是 1 至 24 位数字，前导零会保留。")
        if not name or len(name) > 40 or len(group_name) > 60:
            raise ValueError("请填写姓名（不超过 40 字）和班级（不超过 60 字）。")
        return student_id, name, group_name

    def save_student(self, student_id, name, group_name="未分班"):
        values = self.validate_student(student_id, name, group_name)
        with self.connect(write=True) as connection:
            connection.execute("""INSERT INTO students(student_id,name,group_name) VALUES(?,?,?)
                ON CONFLICT(student_id) DO UPDATE SET name=excluded.name,group_name=excluded.group_name""", values)
            self._audit(connection, "保存学生", student_id=values[0], detail=f"{values[1]} / {values[2]}")

    def set_student_active(self, student_id, active):
        with self.connect(write=True) as connection:
            if not connection.execute("UPDATE students SET active=? WHERE student_id=?", (int(active), student_id)).rowcount:
                raise ValueError("学生不存在。")
            self._audit(connection, "启用学生" if active else "停用学生", student_id=student_id)

    def students(self, active_only=False, group_name=None):
        query, args = "SELECT * FROM students WHERE 1=1", []
        if active_only:
            query += " AND active=1"
        if group_name:
            query += " AND group_name=?"
            args.append(group_name)
        with self.connect() as connection:
            return [dict(row) for row in connection.execute(query + " ORDER BY student_id", args)]

    @classmethod
    def _prepare_import(cls, rows):
        prepared = [cls.validate_student(*row) for row in rows]
        if not prepared:
            raise ValueError("文件没有学生数据。")
        if len({row[0] for row in prepared}) != len(prepared):
            raise ValueError("导入文件包含重复学号，请先修正。")
        return prepared

    def _import_students(self, connection, prepared):
        for row in prepared:
            checkpoint()
            connection.execute("""INSERT INTO students(student_id,name,group_name) VALUES(?,?,?)
                ON CONFLICT(student_id) DO UPDATE SET name=excluded.name,group_name=excluded.group_name""", row)
        self._audit(connection, "批量导入学生", detail=f"共 {len(prepared)} 人，已有学生更新姓名和班级")

    def import_students(self, rows):
        prepared = self._prepare_import(rows)
        with self.connect(write=True) as connection:
            self._import_students(connection, prepared)
        return len(prepared)

    def create_session(self, course_name, group_name, starts_at, ends_at, late_minutes=10, student_ids=None, *,
                       opens_at=None, late_at=None):
        course_name, group_name = course_name.strip(), group_name.strip()
        if not course_name or len(course_name) > 80 or not group_name or len(group_name) > 60:
            raise ValueError("请填写课程与班级，课程不超过 80 字，班级不超过 60 字。")
        times = resolve_session_times(starts_at, ends_at, late_minutes, opens_at=opens_at, late_at=late_at)
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            members = self._session_members(connection, group_name, student_ids)
            return self._insert_session(connection, course_name, group_name, times['starts_at'], times['ends_at'], late_minutes,
                                        members, opens_at=times['opens_at'], late_at=times['late_at'])

    @staticmethod
    def _session_members(connection, group_name, student_ids=None):
        members = [dict(row) for row in connection.execute("SELECT * FROM students WHERE active=1 ORDER BY student_id")
                   if group_name == "全部学生" or row["group_name"] == group_name]
        if student_ids is not None:
            wanted = set(student_ids)
            members = [student for student in members if student["student_id"] in wanted]
            if {student["student_id"] for student in members} != wanted:
                raise ValueError("名单包含不存在、已停用或不属于该班级的学生。")
        if not members:
            raise ValueError("该班级没有启用的学生，请先维护学生名单。")
        return members

    def _insert_session(self, connection, course_name, group_name, starts_at, ends_at, late_minutes, members, *, actor="",
                        opens_at=None, late_at=None):
        times = resolve_session_times(starts_at, ends_at, late_minutes, opens_at=opens_at, late_at=late_at)
        session_id = connection.execute("""INSERT INTO sessions(course_name,group_name,opens_at,starts_at,late_at,ends_at,created_at)
            VALUES(?,?,?,?,?,?,?)""", (course_name, group_name, stamp(times['opens_at']), stamp(times['starts_at']),
                                      stamp(times['late_at']), stamp(times['ends_at']), stamp())).lastrowid
        connection.executemany("INSERT INTO session_students(session_id,student_id,name,group_name) VALUES(?,?,?,?)",
                               [(session_id, member["student_id"], member["name"], member["group_name"]) for member in members])
        self._audit(connection, "创建课次", session_id, detail=f"{course_name} / {group_name} / {len(members)} 人", actor=actor,
                    after={key: stamp(value) for key, value in times.items()})
        return session_id

    def _expire(self, connection, now):
        expired = connection.execute("SELECT id FROM sessions WHERE state='open' AND ends_at<=?", (stamp(now),)).fetchall()
        for row in expired:
            changed = connection.execute("UPDATE sessions SET state='closed' WHERE id=? AND state='open'", (row[0],)).rowcount
            if changed:
                self._audit(connection, EXPIRE_SESSION, row[0], actor_kind="system")

    def open_session(self, session_id, now=None, *, manual_reason=None, actor=""):
        if manual_reason is not None:
            manual_reason = self._reason(manual_reason)
        now = now or datetime.now()
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._expire(connection, now)
            session = connection.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
            if not session or session["state"] != "draft":
                raise ValueError("只能开启尚未开始的课次。")
            if not session["opens_at"] <= stamp(now) < session["ends_at"]:
                raise ValueError("当前时间不在课次的签到时间范围内。")
            if connection.execute("SELECT 1 FROM sessions WHERE state='open'").fetchone():
                raise ValueError("已有课次正在签到，请先结束该课次。")
            if not connection.execute("SELECT 1 FROM session_students WHERE session_id=? AND included=1", (session_id,)).fetchone():
                raise ValueError("课次应到名单为空，请先调整名单或取消课次。")
            connection.execute("UPDATE sessions SET state='open' WHERE id=?", (session_id,))
            self._audit(connection, OPEN_SESSION, session_id,
                        detail=f"人工点名开课；原因：{manual_reason}" if manual_reason is not None else "", actor=actor)

    def open_manual_session(self, session_id, reason, *, actor="", now=None):
        """沿用时段及唯一开放课次规则，教师核实原因后可在没有模型时开课。"""
        return self.open_session(session_id, now, manual_reason=self._reason(reason), actor=actor)

    def close_session(self, session_id):
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            session = self._session(connection, session_id)
            if session["state"] == "closed":
                return False
            if session["state"] != "open":
                raise ValueError("只能结束签到中的课次；待开启课次请使用取消。")
            connection.execute("UPDATE sessions SET state='closed' WHERE id=? AND state='open'", (session_id,))
            self._audit(connection, CLOSE_SESSION, session_id)
            return True

    def cancel_session(self, session_id, reason, actor=""):
        reason = self._reason(reason)
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            session = self._session(connection, session_id)
            if session["state"] == "cancelled":
                return False
            if session["state"] != "draft":
                raise ValueError("只能取消待开启课次；签到中的课次请使用结束。")
            connection.execute("UPDATE sessions SET state='cancelled',cancel_reason=?,cancelled_at=? WHERE id=?",
                               (reason, stamp(), session_id))
            self._audit(connection, CANCEL_SESSION, session_id, detail=reason, actor=actor,
                        before=session, after=self._session(connection, session_id))
            return True

    def active_session(self, now=None):
        now = now or datetime.now()
        with self.connect() as connection:
            self._expire(connection, now)
            row = connection.execute("SELECT * FROM sessions WHERE state='open' AND opens_at<=? AND ends_at>?",
                                     (stamp(now), stamp(now))).fetchone()
            return dict(row) if row else None

    def sessions(self):
        with self.connect() as connection:
            self._expire(connection, datetime.now())
            return [dict(row) for row in connection.execute(f"""SELECT s.*,
                CASE WHEN s.state IN ('open','closed') THEN
                    (SELECT COUNT(*) FROM session_students ss WHERE ss.session_id=s.id AND ss.included=1) ELSE 0 END AS expected,
                (SELECT COUNT(*) FROM session_students ss WHERE ss.session_id=s.id AND ss.included=1) AS roster_count,
                (SELECT COUNT(*) FROM checkins c JOIN session_students ss
                 ON ss.session_id=c.session_id AND ss.student_id=c.student_id
                 WHERE c.session_id=s.id AND ss.included=1 AND c.status IN ({ARRIVED_STATUS_SQL})) AS arrived
                FROM sessions s ORDER BY starts_at DESC,id DESC""")]

    def session(self, session_id):
        with self.connect() as connection:
            return self._session(connection, session_id)

    def record_sign_in(self, student_id, session_id, when=None):
        when = when or datetime.now()
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            session = connection.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
            if not session or session["state"] != "open" or not session["opens_at"] <= stamp(when) < session["ends_at"]:
                raise ValueError("当前课次未开放签到或已结束。")
            member = connection.execute("""SELECT ss.name FROM session_students ss JOIN students s USING(student_id)
                WHERE ss.session_id=? AND ss.student_id=? AND ss.included=1 AND s.active=1""", (session_id, student_id)).fetchone()
            if not member:
                raise ValueError("该学生不在本课次名单中，或已被停用。")
            require_authorized(connection, student_id, when)
            require_authorized(connection, student_id)
            status = "late" if stamp(when) >= session["late_at"] else "present"
            cursor = connection.execute("INSERT OR IGNORE INTO checkins VALUES(?,?,?,?,?,?)",
                                        (session_id, student_id, status, stamp(when), "face", ""))
            evidence = connection.execute("INSERT OR IGNORE INTO face_evidence VALUES(?,?,?,?,?,?)",
                                          (session_id, student_id, stamp(when), "face", stamp(), "captured"))
            if cursor.rowcount:
                self._audit(connection, FACE_CHECKIN, session_id, student_id, STATUS_TEXT[status], actor_kind="face_recognition")
            elif evidence.rowcount:
                self._audit(connection, SAVE_FACE_EVIDENCE, session_id, student_id,
                            "保留已有人工结果", actor_kind="face_recognition")
            record = connection.execute("SELECT * FROM checkins WHERE session_id=? AND student_id=?", (session_id, student_id)).fetchone()
            return dict(record, inserted=bool(cursor.rowcount), name=member["name"])

    def correct_attendance(self, session_id, student_id, status, reason, when=None, *,
                           arrival_time=UNSET, time_unconfirmed=False, actor="", note=UNSET):
        reason = self._reason(reason)
        if status not in {"present", "late", "leave", "absent"}:
            raise ValueError("请选择考勤状态并填写修改原因（1 至 300 字）。")
        if when is not None:
            if arrival_time is not UNSET or time_unconfirmed:
                raise ValueError("请仅填写一种实际到课时间。")
            arrival_time = when
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._expire(connection, datetime.now())
            session = self._session(connection, session_id)
            plan = self._correction_plan(connection, session, student_id, status, reason,
                                         arrival_time, time_unconfirmed, actor, note)
            return self._apply_correction(connection, plan, actor)

    def _correction_plan(self, connection, session, student_id, status, reason,
                         arrival_time=UNSET, time_unconfirmed=False, actor="", note=UNSET):
        if session["state"] not in {"open", "closed"}:
            raise ValueError("只有签到中或已结束的课次可以更正。")
        if len(actor.strip()) > 80:
            raise ValueError("操作署名不超过 80 字。")
        session_id = session["id"]
        member = connection.execute("SELECT * FROM session_students WHERE session_id=? AND student_id=? AND included=1",
                                    (session_id, student_id)).fetchone()
        if not member:
            raise ValueError("该学生不在课次名单中。")
        previous = connection.execute("SELECT * FROM checkins WHERE session_id=? AND student_id=?", (session_id, student_id)).fetchone()
        before = dict(previous) if previous else dict(session_id=session_id, student_id=student_id,
            status="absent" if session["state"] == "closed" else "pending", signed_at=None, method=None, note="")
        original = connection.execute("SELECT * FROM face_evidence WHERE session_id=? AND student_id=?", (session_id, student_id)).fetchone()
        signed_at = before["signed_at"] if previous else (original["signed_at"] if original else None)
        existing_unknown = bool(previous and is_time_unknown(previous) and arrival_time is UNSET)
        if time_unconfirmed:
            if arrival_time is not UNSET and arrival_time is not None:
                raise ValueError("实际到课时间与时间未确认不能同时填写。")
            signed_at = None
        elif arrival_time is not UNSET:
            if not isinstance(arrival_time, datetime) or arrival_time.tzinfo is not None:
                raise ValueError("请填写有效的本地实际到课时间，或明确选择时间未确认。")
            signed_at = stamp(arrival_time)
        if status in ARRIVED_STATUSES and signed_at is None and not time_unconfirmed and not existing_unknown:
            raise ValueError("人工补签须填写实际到课时间，或明确选择时间未确认。")
        current_note = before["note"] if note is UNSET else str(note).strip()
        if len(current_note) > 300:
            raise ValueError("备注不超过 300 字。")
        after = dict(session_id=session_id, student_id=student_id, status=status,
                     signed_at=signed_at, method="manual", note=current_note)
        detail = f"{before['status']} → {status}；原因：{reason}"
        last = connection.execute("SELECT * FROM audit_log WHERE action=? AND session_id=? AND student_id=? ORDER BY id DESC LIMIT 1",
                                  (CORRECT_ATTENDANCE, session_id, student_id)).fetchone()
        unchanged = (before == after and last and last["after_json"] == json.dumps(after, ensure_ascii=False, sort_keys=True)
                     and last["detail"].partition("；原因：")[2] == reason and last["actor_label"] == actor.strip())
        return {"student_id": student_id, "name": member["name"], "member": dict(member),
                "evidence": dict(original) if original else None, "last_correction_id": last["id"] if last else None,
                "before": before, "after": after, "detail": detail, "changed": not unchanged}

    def _apply_correction(self, connection, plan, actor):
        if not plan["changed"]:
            return False
        row = plan["after"]
        connection.execute("""INSERT INTO checkins VALUES(?,?,?,?,?,?) ON CONFLICT(session_id,student_id)
            DO UPDATE SET status=excluded.status,signed_at=excluded.signed_at,method=excluded.method,note=excluded.note""",
            (row["session_id"], row["student_id"], row["status"], row["signed_at"], "manual", row["note"]))
        self._audit(connection, CORRECT_ATTENDANCE, row["session_id"], row["student_id"], plan["detail"],
                    actor=actor, before=plan["before"], after=row)
        return True

    def _manual_attendance_plan(self, connection, session_id, student_ids, status, reason,
                                arrival_time, time_unconfirmed, actor):
        reason = self._reason(reason)
        if status not in {"present", "late", "leave", "absent"}:
            raise ValueError("请选择有效的人工考勤状态。")
        ids = list(student_ids)
        if not ids or any(not isinstance(sid, str) for sid in ids) or len(set(ids)) != len(ids):
            raise ValueError("请选中学生，名单中不能有重复学号。")
        session = self._session(connection, session_id)
        rows = [self._correction_plan(connection, session, sid, status, reason, arrival_time,
                                     time_unconfirmed, actor) for sid in sorted(ids)]
        context = {"session": session, "rows": rows, "actor": actor.strip()}
        token = hashlib.sha256(json.dumps(context, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        return {"session": session, "rows": rows, "token": token,
                "changed": sum(row["changed"] for row in rows), "total": len(rows)}

    def preview_manual_attendance(self, session_id, student_ids, status, reason, *,
                                  arrival_time=UNSET, time_unconfirmed=False, actor=""):
        """预览只读考勤结果；绑定所选人员、时间、原因与原始证据。"""
        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._expire(connection, datetime.now())
            return self._manual_attendance_plan(connection, session_id, student_ids, status, reason,
                                                arrival_time, time_unconfirmed, actor)

    def apply_manual_attendance(self, session_id, student_ids, status, reason, token, *,
                                arrival_time=UNSET, time_unconfirmed=False, actor=""):
        """确认时重新计算预览；所有结果和逐人更正历史在同一事务提交。"""
        student_ids = tuple(student_ids)
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._expire(connection, datetime.now())
            plan = self._manual_attendance_plan(connection, session_id, student_ids, status, reason,
                                                arrival_time, time_unconfirmed, actor)
            if not token or token != plan["token"]:
                raise ValueError("课次、所选名单或考勤已变化，请重新预览后确认。")
            for row in plan["rows"]:
                checkpoint()
                self._apply_correction(connection, row, actor)
            if plan["changed"]:
                self._audit(connection, "批量人工点名", session_id,
                            detail=f"所选 {plan['total']} 人，登记 {plan['changed']} 人；原因：{self._reason(reason)}",
                            actor=actor, after={"student_ids": sorted(student_ids), "status": status,
                                                "changed": plan["changed"]})
            return {"changed": plan["changed"], "unchanged": plan["total"] - plan["changed"]}

    def session_rows(self, session_id, now=None, include_removed=False):
        now = now or datetime.now()
        with self.connect() as connection:
            self._expire(connection, now)
            self._session(connection, session_id)
            return self._session_rows(connection, session_id, now, include_removed)

    @staticmethod
    def _session_rows(connection, session_id, now, include_removed=False):
        query = ATTENDANCE_ROW_SQL + " WHERE ss.session_id=:sid"
        if not include_removed:
            query += " AND ss.included=1"
        return [dict(row) for row in connection.execute(query + " ORDER BY ss.student_id", {"sid":session_id,"now":stamp(now)})]

    def attendance_history(self, session_id, student_id):
        with self.connect() as connection:
            evidence = connection.execute("SELECT * FROM face_evidence WHERE session_id=? AND student_id=?", (session_id, student_id)).fetchone()
            changes = [dict(row) for row in connection.execute("SELECT * FROM audit_log WHERE session_id=? AND student_id=? AND action=? ORDER BY id", (session_id, student_id, CORRECT_ATTENDANCE))]
            return {"evidence": dict(evidence) if evidence else None, "corrections": changes}

    def _roster_plan(self, connection, session_id, refresh, add_ids, remove_ids):
        session = self._session(connection, session_id)
        if session["state"] == "cancelled":
            raise ValueError("已取消课次只能查询，不能调整名单。")
        if refresh and session["state"] != "draft":
            raise ValueError("只有待开启课次可以按班级刷新名单。")
        before = [dict(row) for row in connection.execute("""SELECT ss.student_id,ss.name,ss.group_name,ss.included,
            c.status,c.signed_at,c.method,c.note FROM session_students ss
            LEFT JOIN checkins c ON c.session_id=ss.session_id AND c.student_id=ss.student_id
            WHERE ss.session_id=? ORDER BY ss.student_id""", (session_id,))]
        current = {row["student_id"] for row in before if row["included"]}
        students = {row["student_id"]: dict(row) for row in connection.execute("SELECT * FROM students ORDER BY student_id")}
        additions, removals = set(add_ids), set(remove_ids)
        if additions & removals:
            raise ValueError("同一学生不能同时增补和移出。")
        if refresh:
            if additions or removals:
                raise ValueError("刷新班级名单与手工调整不能同时执行。")
            wanted = {sid for sid, row in students.items() if row["active"] and
                      (session["group_name"] == "全部学生" or row["group_name"] == session["group_name"])}
            additions, removals = wanted-current, current-wanted
        for sid in additions:
            if sid not in students or not students[sid]["active"]:
                raise ValueError(f"增补学号 {sid} 不存在或已停用。")
        if removals-current:
            raise ValueError("只能移出当前应到名单中的学生。")
        additions -= current
        wanted = (current | additions)-removals
        after = [dict(row, included=int(row["student_id"] in wanted)) for row in before]
        known = {row["student_id"] for row in before}
        for sid in sorted(additions-known):
            row = students[sid]
            after.append(dict(student_id=sid,name=row["name"],group_name=row["group_name"],included=1,
                              status=None,signed_at=None,method=None,note=None))
        after.sort(key=lambda row: row["student_id"])
        # 预览同时绑定课次状态、已有考勤和当前档案，防止确认时悄悄应用另一份名单。
        context = {"state":session["state"],"before":before,"after":after,
                   "candidates":[students[sid] for sid in sorted(additions)],"refresh":bool(refresh)}
        token = hashlib.sha256(json.dumps(context,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
        return {"session_id":session_id,"state":session["state"],"before":before,"after":after,
                "added":sorted(additions),"removed":sorted(removals),"token":token,
                "recorded_removals":[row["student_id"] for row in before if row["student_id"] in removals and row["method"]]}

    def preview_roster_change(self, session_id, *, refresh=False, add_ids=(), remove_ids=()):
        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._expire(connection, datetime.now())
            return self._roster_plan(connection, session_id, refresh, add_ids, remove_ids)

    def apply_roster_change(self, session_id, token, reason, *, refresh=False, add_ids=(), remove_ids=(), actor=""):
        reason = self._reason(reason)
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._expire(connection, datetime.now())
            plan = self._roster_plan(connection, session_id, refresh, add_ids, remove_ids)
            if not token or token != plan["token"]:
                raise ValueError("课次、名单或考勤已变化，请重新预览后确认。")
            if not plan["added"] and not plan["removed"]:
                return False
            for row in plan["after"]:
                connection.execute("""INSERT INTO session_students(session_id,student_id,name,group_name,included)
                    VALUES(?,?,?,?,?) ON CONFLICT(session_id,student_id) DO UPDATE SET included=excluded.included""",
                    (session_id,row["student_id"],row["name"],row["group_name"],row["included"]))
            self._audit(connection, CHANGE_ROSTER, session_id,
                        detail=f"增补 {', '.join(plan['added']) or '无'}；移出 {', '.join(plan['removed']) or '无'}；原因：{reason}",
                        actor=actor, before=plan["before"], after=plan["after"])
            return True

    def audit_rows(self, session_id=None):
        with self.connect() as connection:
            query = "SELECT * FROM audit_log"
            args = ()
            if session_id is not None:
                query += " WHERE session_id=?"
                args = (session_id,)
            return [dict(row) for row in connection.execute(query + " ORDER BY id DESC LIMIT 500", args)]

    def semester_dates(self, now=None):
        with self.connect() as connection:
            value = self._get(connection, "semester_dates")
        if not value:
            return default_semester_dates(now)
        try:
            saved = json.loads(value)
            return validate_semester_dates(saved["starts_on"], saved["ends_on"])
        except (ValueError, TypeError, KeyError):
            raise ValueError("学期设置无效，请重新设置学期起止日期或恢复默认。") from None

    def set_semester_dates(self, starts_on, ends_on):
        start, end = validate_semester_dates(starts_on, ends_on)
        with self.connect(write=True) as connection:
            self._set(connection, "semester_dates", json.dumps({"starts_on": start.isoformat(), "ends_on": end.isoformat()}))
            self._audit(connection, "设置学期", detail=f"{start} 至 {end}")

    def reset_semester_dates(self):
        with self.connect(write=True) as connection:
            self._set(connection, "semester_dates", "")
            self._audit(connection, "设置学期", detail="恢复默认学期划分")

    def attendance_records(self, starts_at, ends_at, status="all", session_id=None, now=None):
        if starts_at >= ends_at:
            raise ValueError("考勤查询结束时间须晚于开始时间。")
        if status not in {"all", "absent"}:
            raise ValueError("考勤筛选仅支持全部或缺勤。")
        now = now or datetime.now()
        params = {"start": stamp(starts_at), "end": stamp(ends_at), "now": stamp(now)}
        query = "SELECT * FROM (" + ATTENDANCE_ROW_SQL + """
            WHERE ss.included=1 AND s.state IN ('open','closed') AND s.starts_at>=:start AND s.starts_at<:end
        ) WHERE 1=1"""
        if status == "absent":
            query += " AND status='absent'"
        if session_id is not None:
            query += " AND session_id=:sid"
            params["sid"] = int(session_id)
        query += " ORDER BY session_starts_at DESC,session_id DESC,student_id"
        with self.connect() as connection:
            self._expire(connection, now)
            return [dict(row) for row in connection.execute(query, params)]

    def admin_exists(self):
        with self.connect() as connection:
            return bool(self._get(connection, "admin_auth"))

    def setup_admin(self, password):
        if len(password) < 8 or len(password) > 128 or password.isspace():
            raise ValueError("密码应为 8 至 128 个字符。")
        salt = secrets.token_hex(16)
        hashed = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), PASSWORD_ITERATIONS).hex()
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            if self._get(connection, "admin_auth"):
                raise ValueError("管理员已设置，请使用现有密码登录。")
            self._set(connection, "admin_auth", json.dumps({"algorithm":"pbkdf2_sha256", "iterations":PASSWORD_ITERATIONS,
                      "salt": salt, "hash": hashed, "failures": 0, "locked_until": ""}))
            self._audit(connection, "设置管理员密码")

    @staticmethod
    def _auth_status(auth, now):
        auth = auth or {}
        failures = auth.get("failures", 0)
        if type(failures) is not int or failures < 0:
            raise ValidationError("管理口令信息无效，请按离线重置流程处理。")
        locked_until = auth.get("locked_until", "")
        try:
            retry_after = max(0, ceil((datetime.fromisoformat(locked_until)-now).total_seconds())) if locked_until else 0
        except (ValueError, TypeError) as error:
            raise ValidationError("管理口令信息无效，请按离线重置流程处理。") from error
        return {"configured": bool(auth), "failures": failures,
                "attempts_remaining": (0 if retry_after else max(1, 5-failures)) if auth else 0,
                "locked_until": locked_until, "retry_after_seconds": retry_after,
                "next_lock_seconds": min(3600, 60 * 2**min(max(failures-4, 0), 6))}

    def admin_auth_status(self, now=None) -> dict:
        """只读返回登录次数与锁定时间，不返回口令摘要，也不重置失败计数。"""
        with self.connect() as connection:
            value = self._get(connection, "admin_auth")
            return self._auth_status(json.loads(value) if value else None, now or datetime.now())

    def authenticate_with_status(self, password, now=None) -> dict:
        """在同一事务验证并返回本次结果、剩余次数及锁定状态。"""
        now = now or datetime.now()
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            value = self._get(connection, "admin_auth")
            if not value:
                return dict(self._auth_status(None, now), authenticated=False, attempted=False)
            auth = json.loads(value)
            status = self._auth_status(auth, now)
            if status["retry_after_seconds"]:
                return dict(status, authenticated=False, attempted=False)
            iterations = auth.get("iterations", LEGACY_PASSWORD_ITERATIONS)
            if not isinstance(iterations,int) or not 1 <= iterations <= 2_000_000:
                raise ValidationError("管理口令信息无效，请按离线重置流程处理。")
            hashed = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(auth["salt"]), iterations).hex()
            valid = hmac.compare_digest(auth["hash"], hashed)
            auth["failures"] = 0 if valid else auth["failures"] + 1
            seconds = min(3600, 60 * 2**min(max(auth["failures"]-5,0),6))
            auth["locked_until"] = stamp(now + timedelta(seconds=seconds)) if not valid and auth["failures"] >= 5 else ""
            if valid and iterations < PASSWORD_ITERATIONS:
                auth["salt"] = secrets.token_hex(16)
                auth["hash"] = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(auth["salt"]), PASSWORD_ITERATIONS).hex()
                auth["iterations"] = PASSWORD_ITERATIONS
                auth["algorithm"] = "pbkdf2_sha256"
                self._audit(connection,"升级管理口令参数")
            self._set(connection, "admin_auth", json.dumps(auth))
            return dict(self._auth_status(auth, now), authenticated=valid, attempted=True)

    def authenticate(self, password, now=None):
        """保留布尔验证接口；锁定期间拒绝验证并给出剩余秒数。"""
        result = self.authenticate_with_status(password, now)
        if result["retry_after_seconds"] and not result["attempted"]:
            raise ValidationError(f"登录尝试过多，还需等待 {result['retry_after_seconds']} 秒后重试。")
        return result["authenticated"]

    def change_password(self, old_password, new_password):
        """验证旧密码，原子替换摘要，不保存或记录明文。"""
        if not 8 <= len(new_password) <= 128 or new_password.isspace():
            raise ValidationError("新密码应为 8 至 128 个字符。")
        if not self.authenticate(old_password):
            raise ValidationError("原管理密码不正确。")
        with self.connect(write=True) as connection:
            connection.execute("BEGIN IMMEDIATE")
            # 防止另一个管理员在验证与修改之间已更换密码。
            auth=json.loads(self._get(connection,"admin_auth"))
            hashed=hashlib.pbkdf2_hmac("sha256",old_password.encode(),bytes.fromhex(auth["salt"]),
                                     auth.get("iterations",LEGACY_PASSWORD_ITERATIONS)).hex()
            if not hmac.compare_digest(auth["hash"],hashed):
                raise ValidationError("管理密码已变更，请重新验证。")
            salt=secrets.token_hex(16)
            value={"algorithm":"pbkdf2_sha256","iterations":PASSWORD_ITERATIONS,"salt":salt,
                   "hash":hashlib.pbkdf2_hmac("sha256",new_password.encode(),bytes.fromhex(salt),PASSWORD_ITERATIONS).hex(),
                   "failures":0,"locked_until":""}
            self._set(connection,"admin_auth",json.dumps(value))
            self._audit(connection,"修改管理密码")

    def camera_index(self, value=None):
        with self.connect(write=value is not None) as connection:
            if value is not None:
                if not 0 <= int(value) <= 9:
                    raise ValueError("摄像头编号应为 0 至 9。")
                self._set(connection, "camera_index", str(int(value)))
            return int(self._get(connection, "camera_index", "0"))

    def export_excel(self, path, session_id):
        session = self.session(session_id)
        if session["state"] not in {"open", "closed"}:
            raise ValueError("待开启或已取消课次不参与考勤导出。")
        rows = self.session_rows(session_id)
        values = [[i, *attendance_base_values(row),
                   excel_time(row["original_signed_at"]), arrival_time_note(row),
                   *[datetime.fromisoformat(session[key]) for key in ('opens_at', 'starts_at', 'late_at', 'ends_at')]]
                  for i, row in enumerate(rows, 1)]
        _write_excel(path, f"{session['course_name']} · {session['group_name']}",
                     f"开放签到：{session['opens_at'].replace('T',' ')}；正式上课：{session['starts_at'].replace('T',' ')}；迟到起点：{session['late_at'].replace('T',' ')}；关闭签到：{session['ends_at'].replace('T',' ')}",
                     ["序号", "学号", "姓名", "班级", "状态", "实际到课时间", "记录方式", "备注", "原始人脸签到时间", "到课时间说明",
                      "开放签到时间", "正式上课时间", "迟到起点", "关闭签到时间"],
                     values, widths=[9,27,16,22,13,24,16,45,24,20,24,24,24,24], date_columns=(6,9,11,12,13,14))

    def export_attendance(self, path, starts_at, ends_at, status="all", session_id=None, title="考勤记录", now=None):
        rows = self.attendance_records(starts_at, ends_at, status, session_id, now)
        values = [[i, *attendance_base_values(row),
                   row["course_name"], row["session_id"], datetime.fromisoformat(row["session_starts_at"]),
                   excel_time(row["original_signed_at"]), arrival_time_note(row),
                   *[datetime.fromisoformat(row[key]) for key in ('session_opens_at', 'session_late_at', 'session_ends_at')]]
                  for i, row in enumerate(rows, 1)]
        subtitle = f"正式上课日期：{starts_at:%Y-%m-%d} 至 {ends_at-timedelta(days=1):%Y-%m-%d}；{'仅缺勤' if status=='absent' else '全部状态'}"
        if session_id is not None:
            subtitle += f"；课次 #{session_id}"
        _write_excel(path, title, subtitle,
                     ["序号", "学号", "姓名", "班级", "状态", "实际到课时间", "记录方式", "备注", "课程", "课次编号", "正式上课时间", "原始人脸签到时间", "到课时间说明",
                      "开放签到时间", "迟到起点", "关闭签到时间"],
                     values, widths=[9,27,16,22,13,24,16,45,28,12,24,24,20,24,24,24], date_columns=(6,11,12,14,15,16))

    def backup_archive(self, path):
        """数据库、模型和样本统一属于当前数据库目录，锁也取自该目录。"""
        from backup_service import create_backup
        from resource_lock import ResourceLease
        data = self.path.resolve().parent
        with ResourceLease(data / '.samples.lock', '正在采集或训练，请完成后再备份。'):
            info = create_backup(self, path, model=data / 'face_model.yml',
                                 manifest=data / 'model_manifest.json', dataset=data / 'dataset')
            self.record_backup_success()
        return info


def _write_excel(path, title, subtitle, headers, rows, widths=None, date_columns=(6,)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("考勤导出文件必须使用 .xlsx 扩展名。")
    require_disk_space(path.parent)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "考勤记录"
    sheet.append([title]); sheet.append([subtitle]); sheet.append([]); sheet.append(headers)
    for row in rows:
        sheet.append(row)
    # 标题也来自用户输入，和姓名、备注一样按文本写入。
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{sheet.cell(4,len(headers)).column_letter}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="17243C")
    sheet.row_dimensions[1].height = 34
    for cell in sheet[4]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3267E3")
    for column, width in enumerate(widths or [9,27,16,22,13,24,16,45],1):
        sheet.column_dimensions[sheet.cell(4,column).column_letter].width = width
    for row in sheet.iter_rows(min_row=5):
        sheet.row_dimensions[row[0].row].height = 26
        for cell in row:
            if isinstance(cell.value,str):
                cell.data_type = "s"
            cell.font = Font(name="Microsoft YaHei", size=11, color="334155")
            cell.alignment = Alignment(vertical="center")
            if cell.row % 2:
                cell.fill = PatternFill("solid", fgColor="F3F6FC")
        row[1].number_format = "@"
        for column in date_columns:
            if len(row) >= column:
                row[column-1].number_format = "yyyy-mm-dd hh:mm:ss"
    path = Path(path)
    temporary = path.with_name(f".{path.stem}.tmp.xlsx")
    try:
        workbook.save(temporary)
        checkpoint()
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def read_student_import(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="gb18030")
        matrix = list(csv.reader(text.splitlines()))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            matrix = list(workbook.active.values)
        finally:
            workbook.close()
    else:
        raise ValueError("请选择 .xlsx 或 .csv 名单文件。")
    if not matrix:
        raise ValueError("名单文件为空。")
    header = [str(cell or "").strip() for cell in matrix[0]]
    if not {"学号", "姓名"}.issubset(header):
        raise ValueError("第一行需包含“学号”“姓名”列，可选“班级”列。")
    rows = []
    for number, raw in enumerate(matrix[1:], 2):
        if not any(value is not None and str(value).strip() for value in raw):
            continue
        def get(column, default=""):
            if column not in header:
                return default
            index = header.index(column)
            return raw[index] if index < len(raw) and raw[index] is not None else default
        student_id = get("学号")
        if isinstance(student_id, (int, float)):
            if student_id != int(student_id) or len(str(int(student_id))) > 15:
                raise ValueError(f"第 {number} 行学号可能丢失精度，请在 Excel 中按文本保存学号。")
            student_id = str(int(student_id))
        try:
            rows.append(AttendanceStore.validate_student(student_id, get("姓名"), get("班级", "未分班")))
        except ValueError as error:
            raise ValueError(f"第 {number} 行：{error}") from error
    return rows
