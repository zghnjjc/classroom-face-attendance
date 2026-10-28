"""发行包自检：使用临时库和合成数据，不使用真实摄像头或正式数据。"""

import json
import multiprocessing
import queue
import sqlite3
import sys
import tempfile
import zipfile
from datetime import datetime, timedelta
from contextlib import closing
from pathlib import Path


def _echo_worker(pipe):
    pipe.send("frozen-worker-ready")
    pipe.close()


def run_self_test(report):
    result = {"bundled": bool(getattr(sys, "frozen", False)), "python": sys.version.split()[0]}
    try:
        import cv2
        import numpy as np
        import openpyxl
        import tkinter as tk
        from classroom_attendance.core.app_config import CASCADE_PATH, DATA_DIR, PROGRAM_DIR, SEED_PATH
        from classroom_attendance.faces.face_utils import load_classifier, save_gray, read_gray
        root = tk.Tk()
        root.withdraw()
        result["tk"] = str(root.tk.call("info", "patchlevel"))
        root.destroy()
        result["opencv"] = cv2.__version__
        result["face_module"] = hasattr(cv2, "face")
        result["cascade"] = not load_classifier().empty()
        result["resources"] = CASCADE_PATH.is_file() and SEED_PATH.is_file()
        result["separate_data"] = DATA_DIR != PROGRAM_DIR and not DATA_DIR.is_relative_to(PROGRAM_DIR)
        with tempfile.TemporaryDirectory(prefix="classroom-package-check-") as temporary:
            folder = Path(temporary)
            with closing(sqlite3.connect(folder / "test.db")) as connection:
                result["sqlite"] = connection.execute("SELECT 1").fetchone()[0] == 1
            image = folder / "合成图像.png"
            save_gray(image, np.full((160,160), 100, dtype=np.uint8))
            result["image_io"] = read_gray(image).shape == (160,160)
            cv2.face.LBPHFaceRecognizer_create().train([read_gray(image)], np.asarray([1], dtype=np.int32))
            result["lbph"] = True
            book = openpyxl.Workbook()
            book.save(folder / "test.xlsx")
            book.close()
            from classroom_attendance.storage.attendance_store import AttendanceStore
            from classroom_attendance.storage.backup_service import create_backup, inspect_backup, restore_backup, database_summary
            from classroom_attendance.faces.train import train_model, manifest_path
            source = folder / 'source'
            store = AttendanceStore(source / 'attendance.db', seed_path=None)
            store.save_student('001','合成测试','测试班')
            now = datetime.now().replace(microsecond=0)
            sid = store.create_session('合成人工点名', '测试班', now+timedelta(minutes=5), now+timedelta(hours=1),
                                       opens_at=now-timedelta(minutes=1), late_at=now+timedelta(minutes=10))
            store.open_manual_session(sid, '发行包人工入口自检')
            plan = store.preview_manual_attendance(sid, ['001'], 'present', '合成数据自检', time_unconfirmed=True)
            store.apply_manual_attendance(sid, ['001'], 'present', '合成数据自检', plan['token'], time_unconfirmed=True)
            # 冻结包中也实际构建新窗口，验证动态模块和 Tk 控件均已收集。
            from classroom_attendance.attendance.manual_attendance import ManualAttendanceDialog
            class Owner:
                async_io = False
                def say(self, message): pass
                def refresh_sessions(self): pass
            owner = Owner()
            owner.root, owner.store = tk.Tk(), store
            owner.root.withdraw()
            dialog = ManualAttendanceDialog(owner, sid)
            owner.root.update_idletasks()
            result['manual_attendance'] = (len(dialog.roster_tree.get_children()) == 1
                                           and store.session_rows(sid)[0]['status'] == 'present')
            dialog.close()
            from classroom_attendance.ui.student_app import StudentApp
            student_root = tk.Toplevel(owner.root)
            student_root.withdraw()
            student = StudentApp(student_root, store, async_io=False)
            result['session_times'] = (student.current_session['id'] == sid and store.active_session(now)['id'] == sid
                                        and all(title in student.details.cget('text') for title in ('开放签到', '正式上课', '迟到起点', '关闭签到')))
            student_root.update_idletasks()
            student.close()
            from classroom_attendance.courses.course_schedule_ui import CourseScheduleDialog
            first = (now + timedelta(days=7)).date()
            last = first + timedelta(days=7)
            definition = dict(name='合成周期模板', course_name='合成周课', group_name='测试班', start_time='08:00', end_time='09:30',
                              opens_time='07:45', late_time='08:10', opens_previous_day=False, late_next_day=False,
                              late_minutes=10, weekdays=[first.weekday()], interval_weeks=1, next_day=False, student_ids=None)
            template = store.save_course_template(definition)
            plan = store.preview_course_schedule(definition, first, last)
            course_ids = store.create_course_schedule(definition, first, last, plan['token'])
            owner.course = tk.StringVar(master=owner.root, value='合成周课')
            owner.session_group = tk.StringVar(master=owner.root, value='测试班')
            owner.grace = tk.StringVar(master=owner.root, value='10')
            course_dialog = CourseScheduleDialog(owner)
            owner.root.update_idletasks()
            result['course_schedule'] = len(course_ids) == 2 and len(course_dialog.templates) == 1
            course_dialog.close()
            owner.root.destroy()
            from classroom_attendance.ui.admin_app import AdminApp
            store.close_session(sid)
            dashboard_root = tk.Tk()
            dashboard_root.withdraw()
            admin = AdminApp(dashboard_root, store, authenticated=True, async_io=False)
            admin.show_review(sid)
            review = admin.review_dialog
            review.confirmed.set(True)
            review.save()
            result['teacher_dashboard'] = (admin.notebook.select() == str(admin.pages['home']) and review.saved
                                           and admin.dashboard.snapshot['pending_review_count'] == 0)
            report_start = datetime.fromisoformat(store.session(sid)['starts_at']).date().isoformat()
            report_end = report_start
            admin.reports.starts_on.set(report_start)
            admin.reports.ends_on.set(report_end)
            admin.reports.refresh()
            attendance_report = admin.reports.snapshot
            report_path = folder / '合成考勤统计.xlsx'
            store.export_attendance_report(report_path, attendance_report)
            report_book = openpyxl.load_workbook(report_path)
            try:
                result['attendance_reports'] = (attendance_report['totals']['attendance_rate'] == 1
                                                 and len(admin.reports.matrix_tree.get_children()) == 1
                                                 and len(report_book.sheetnames) == 5
                                                 and report_book['学生学期汇总']['A5'].value == '001'
                                                 and report_book['考勤矩阵']['D5'].value == '已签到')
            finally:
                report_book.close()
            from classroom_attendance.maintenance.maintenance_tools import MaintenanceService
            from classroom_attendance.maintenance.maintenance_ui import CameraDiagnosticDialog
            maintenance = MaintenanceService(store)
            admin.maintenance.training.set('240')
            admin.maintenance.export_seconds.set('90')
            admin.maintenance.save()
            diagnostic_path = folder / '合成脱敏诊断.zip'
            diagnostic = maintenance.export_diagnostics(diagnostic_path)
            with zipfile.ZipFile(diagnostic_path) as package:
                result['maintenance_tools'] = (set(package.namelist()) == {'diagnostics.json', '说明.txt'}
                                               and diagnostic['record_counts']['students'] == 1
                                               and diagnostic['data_disk']['free_bytes'] > 0
                                               and maintenance.timeouts() == {'training': 240, 'export': 90})
            class ProbePreview:
                def __init__(self, index, lock):
                    self.events = queue.Queue()
                def start(self): pass
                def stop(self): pass
                def take_frame(self): return None
            camera_check = CameraDiagnosticDialog(admin, worker_factory=ProbePreview)
            camera_check.start()
            camera_check.worker.events.put(('error', '合成断开'))
            camera_check.worker.events.put(('stopped', None))
            camera_check.poll()
            released = camera_check.worker is None and '设备已释放' in camera_check.message.get()
            camera_check.start()
            result['camera_retry'] = released and camera_check.worker is not None
            camera_check.worker.events.put(('stopped', None))
            camera_check.poll()
            camera_check.close()
            from classroom_attendance.faces.face_lifecycle import FaceLifecycleService
            lifecycle = FaceLifecycleService(store)
            lifecycle.publish_notice('synthetic-probe-v1', '合成自检告知：仅用于发行包测试，包含人脸资料、明确同意、保存期限、访问与撤回渠道，以及人工点名替代方式。', confirmed=True)
            authorization = lifecycle.authorization_plan('001')
            lifecycle.grant('001', authorization['token'], (now+timedelta(days=30)).date().isoformat(),
                            '本人明确同意', 'synthetic-probe-reference', confirmed=True)
            admin.lifecycle.refresh()
            result['face_lifecycle'] = lifecycle.eligible_ids() == {'001'} and len(admin.lifecycle.roster.get_children()) == 1
            admin.close()
            from classroom_attendance.onboarding.first_use import FirstUseService
            from classroom_attendance.onboarding.first_use_ui import FirstUseWizard
            setup_service = FirstUseService(store)
            setup_service.prepare(manual=True)
            store.setup_admin('Synthetic-probe-password-2026')
            setup_root = tk.Tk()
            setup_root.withdraw()
            targets = []
            setup = FirstUseWizard(setup_root, store, targets.append, async_io=False)
            setup.skip()
            setup_root.update_idletasks()
            setup.close()
            setup_root = tk.Tk()
            setup_root.withdraw()
            setup = FirstUseWizard(setup_root, store, targets.append, async_io=False)
            resumed_roster = setup.page == 'roster'
            setup.next()
            setup.finish('sessions')
            setup_root.update_idletasks()
            setup_root.destroy()
            result['first_use'] = (resumed_roster and targets == ['sessions'] and setup_service.prepare()['completed']
                                   and setup_service.snapshot()['state']['camera'] == 'skipped')
            dataset = source / 'dataset'
            dataset.mkdir()
            save_gray(dataset / '001.0.jpg',np.full((160,160),100,dtype=np.uint8))
            model = source / 'face_model.yml'
            manifest = source / 'model_manifest.json'
            train_model(dataset,model,{'001'},_lease=False,store=store)
            manifest_path(model).replace(manifest)
            from classroom_attendance.storage.automatic_backup import AutomaticBackup
            automatic = AutomaticBackup(store)
            automatic.save_settings(enabled=True, time='00:00', keep=1, directory=str(folder/'automatic'))
            automatic_result = automatic.run_due(force=True)
            automatic_folder = automatic.folder(automatic.settings())
            automatic_archives = list(automatic_folder.glob('auto_*.zip'))
            automatic_summary = inspect_backup(automatic_archives[0])
            result['automatic_backup'] = (automatic_result['outcome'] == 'success' and automatic_summary['model_ready']
                                          and automatic_summary['students'] == 1
                                          and automatic.run_due()['outcome'] == 'idle')
            archive = folder / 'backup.zip'
            create_backup(store,archive,model=model,manifest=manifest,dataset=dataset)
            summary = inspect_backup(archive)
            restored = restore_backup(archive,folder / 'restored',expected_sha256=summary['archive_sha256'])
            result['backup_restore'] = (restored['source_model_ready'] and restored['authorization_review_required']
                                         and not restored['model_ready'] and database_summary(folder / 'restored' / 'attendance.db')['students'] == 1)
            restored_store = AttendanceStore(folder / 'restored' / 'attendance.db', seed_path=None)
            result['manual_attendance'] = result['manual_attendance'] and restored_store.session_rows(sid)[0]['status'] == 'present'
            result['course_schedule'] = result['course_schedule'] and restored_store.course_templates() == [template] and restored_store.session(course_ids[0])['state'] == 'draft'
            result['teacher_dashboard'] = result['teacher_dashboard'] and restored_store.preview_session_review(sid)['reviewed']
            result['automatic_backup'] = result['automatic_backup'] and AutomaticBackup(restored_store).settings().get('relocated') is True
            result['attendance_reports'] = (result['attendance_reports']
                                             and restored_store.attendance_report(report_start, report_end, now=now)['totals']
                                             == attendance_report['totals'])
            result['first_use'] = result['first_use'] and FirstUseService(restored_store).prepare()['completed']
            result['session_times'] = (result['session_times'] and all(restored_store.session(sid)[key] == store.session(sid)[key]
                                                                      for key in ('opens_at', 'starts_at', 'late_at', 'ends_at')))
            result['maintenance_tools'] = result['maintenance_tools'] and MaintenanceService(restored_store).timeouts() == {'training': 240, 'export': 90}
            restored_lifecycle = FaceLifecycleService(restored_store)
            result['face_lifecycle'] = result['face_lifecycle'] and not restored_lifecycle.eligible_ids()
            reviewed = restored_lifecycle.authorization_plan('001')
            restored_lifecycle.grant('001', reviewed['token'], (now+timedelta(days=30)).date().isoformat(),
                                     '本人明确同意', 'synthetic-restored-review', confirmed=True)
            result['face_lifecycle'] = result['face_lifecycle'] and restored_lifecycle.eligible_ids() == {'001'}
            restored_lifecycle.revoke('001', '合成自检撤回')
            result['face_lifecycle'] = result['face_lifecycle'] and not restored_lifecycle.eligible_ids()
            cleanup = restored_lifecycle.preview_cleanup()
            removed = restored_lifecycle.apply_cleanup('Synthetic-probe-password-2026', cleanup['token'], '合成自检清理')
            result['face_lifecycle'] = (result['face_lifecycle'] and removed == 1
                                       and not (folder/'restored'/'face_model.yml').exists()
                                       and restored_store.session_rows(sid)[0]['status'] == 'present')
            from classroom_attendance.storage.backup_service import SnapshotSource, database_schema
            legacy_store = AttendanceStore(folder/'legacy'/'attendance.db', seed_path=None)
            legacy_store.save_student('001', '合成旧库', '测试班')
            legacy_sid = legacy_store.create_session('合成旧课次', '测试班', now, now+timedelta(hours=1))
            with legacy_store.connect(write=True) as connection:
                connection.execute('ALTER TABLE sessions DROP COLUMN opens_at')
                legacy_store._set(connection, 'schema_version', '3')
            legacy_archive = folder/'legacy-schema3.zip'
            create_backup(SnapshotSource(legacy_store.path), legacy_archive, model=folder/'none.yml',
                          manifest=folder/'none.json', dataset=folder/'no-samples')
            legacy_summary = inspect_backup(legacy_archive)
            legacy_target = folder/'legacy-restored'
            restore_backup(legacy_archive, legacy_target, expected_sha256=legacy_summary['archive_sha256'])
            upgraded = AttendanceStore(legacy_target/'attendance.db', seed_path=None)
            result['legacy_schema3_restore'] = (legacy_summary['schema_version'] == 3 and database_schema(upgraded.path) == 5
                                                and upgraded.session(legacy_sid)['opens_at'] == now.isoformat())
        context = multiprocessing.get_context("spawn")
        parent, child = context.Pipe()
        process = context.Process(target=_echo_worker, args=(child,), daemon=True)
        process.start()
        child.close()
        try:
            result["worker_spawn"] = bool(parent.poll(15) and parent.recv() == "frozen-worker-ready")
        finally:
            process.join(2)
            if process.is_alive():
                process.terminate()
                process.join(2)
            parent.close()
        result["ok"] = all(result.get(key) is True for key in
                           ("face_module", "cascade", "resources", "separate_data", "sqlite", "image_io", "lbph", "manual_attendance", "course_schedule", "teacher_dashboard", "automatic_backup", "attendance_reports", "first_use", "session_times", "maintenance_tools", "camera_retry", "face_lifecycle", "legacy_schema3_restore", "backup_restore", "worker_spawn"))
    except Exception as error:
        result.update(ok=False, error_type=type(error).__name__)
    Path(report).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if result["ok"] else 1
