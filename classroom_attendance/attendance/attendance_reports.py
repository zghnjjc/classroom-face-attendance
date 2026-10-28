"""基于课次名单快照生成学期汇总和考勤矩阵。"""

import tempfile
from collections import Counter
from datetime import datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from classroom_attendance.attendance.attendance_filters import validate_semester_dates
from classroom_attendance.core.task_runtime import checkpoint
from classroom_attendance.attendance.attendance_fields import STATUS_TEXT, ARRIVED_STATUSES, is_time_unknown, arrival_time_note, method_text


REPORT_STATUSES = {status: STATUS_TEXT[status] for status in (*ARRIVED_STATUSES, 'leave', 'absent')}
REPORT_RULES = (
    '按正式上课日期归属，范围首尾当天均包含；前日开放或跨日签到不改写归属日期。',
    '只统计已结束课次（含已到关闭时间的签到中课次）；进行中、待开启、已取消课次不计入。',
    '应到人次以各课次仍纳入的名单快照为准；移出名单不计入，停用或转班不改写历史。',
    '出勤率 =（已签到人次 + 迟到人次）÷ 应到人次；迟到计出勤，请假保留在分母中。',
    '准时率 = 已签到人次 ÷ 应到人次；应到为 0 时两种比率均显示“—”。汇总按人次相除，不平均各行百分比。',
    '采用当前有效考勤结果；人工登记时间未确认仍按登记状态计数，首次人脸证据不覆盖更正结果。',
    '班级筛选采用学生在课次中的班级快照；学生汇总姓名/班级采用所选范围最近一次快照。',
    '课程汇总按课程名称与课次班级合并，同名同班课程合并；矩阵每列带独立课次编号，“—”表示不在该次应到名单。',
    '查询生成一致快照；导出使用页面最近生成的同一快照。更正或名单调整后请重新生成。比率显示四舍五入至两位小数。',
)
SUMMARY_FIELDS = ('expected', 'present', 'late', 'leave', 'absent', 'arrived', 'time_unconfirmed',
                  'attendance_rate', 'on_time_rate')
SUMMARY_HEADERS = ('应到人次', '已签到', '迟到', '请假', '缺勤', '已到人次', '到课时间未确认',
                   '出勤率（含迟到）', '准时率')


def summarize(rows):
    counts = Counter(row['status'] for row in rows)
    if counts.keys() - REPORT_STATUSES.keys():
        raise ValueError('考勤结果包含无法统计的状态，请先核对课次记录。')
    expected = len(rows)
    arrived = sum(counts[status] for status in ARRIVED_STATUSES)
    return dict(expected=expected, **{key: counts[key] for key in REPORT_STATUSES}, arrived=arrived,
                time_unconfirmed=sum(is_time_unknown(row) for row in rows),
                attendance_rate=arrived / expected if expected else None,
                on_time_rate=counts['present'] / expected if expected else None)


def rate_text(value):
    if value is None:
        return '—'
    percent = (Decimal(str(value)) * 100).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    return f'{percent:.2f}%'


class AttendanceReportStore:
    def attendance_report(self, starts_on, ends_on, *, course_name=None, group_name=None, now=None):
        start, last = validate_semester_dates(starts_on, ends_on)
        for value in (course_name, group_name):
            if value is not None and (not isinstance(value, str) or not value.strip()):
                raise ValueError('课程或班级筛选无效，请重新选择。')
        now = (now or datetime.now()).replace(microsecond=0)
        params = {'start': datetime.combine(start, time.min).isoformat(),
                  'end': datetime.combine(last + timedelta(days=1), time.min).isoformat(),
                  'now': now.isoformat()}
        scope = "s.starts_at>=:start AND s.starts_at<:end AND (s.state='closed' OR (s.state='open' AND s.ends_at<=:now))"
        with self.connect() as connection:
            connection.execute('BEGIN')
            sessions = [dict(row) for row in connection.execute(
                'SELECT s.* FROM sessions s WHERE ' + scope + ' ORDER BY s.starts_at,s.id', params)]
            rows = []
            for raw in connection.execute(self.attendance_row_sql + ' WHERE ss.included=1 AND ' + scope
                                          + ' ORDER BY s.starts_at,s.id,ss.student_id', params):
                checkpoint()
                rows.append(dict(raw))
        options = {'courses': sorted({session['course_name'] for session in sessions}),
                   'groups': sorted({row['group_name'] for row in rows})}
        rows = [row for row in rows if (course_name is None or row['course_name'] == course_name)
                and (group_name is None or row['group_name'] == group_name)]
        matched_sessions = {row['session_id'] for row in rows}
        sessions = [session for session in sessions if (course_name is None or session['course_name'] == course_name)
                    and (group_name is None or session['id'] in matched_sessions)]
        students, courses = {}, {}
        for session in sessions:
            key = (session['course_name'], session['group_name'])
            course = courses.setdefault(key, dict(course_name=key[0], group_name=key[1], session_ids=[], rows=[]))
            course['session_ids'].append(session['id'])
        for row in rows:
            checkpoint()
            student = students.setdefault(row['student_id'], dict(student_id=row['student_id'], rows=[], cells={}))
            student.update(name=row['name'], group_name=row['group_name'])
            student['rows'].append(row)
            student['cells'][row['session_id']] = row['status']
            courses[(row['course_name'], row['session_group_name'])]['rows'].append(row)
        for group in [*students.values(), *courses.values()]:
            checkpoint()
            group.update(summarize(group.pop('rows')))
        student_rows = sorted(students.values(), key=lambda row: row['student_id'])
        course_rows = sorted(courses.values(), key=lambda row: (row['course_name'], row['group_name']))
        return {'starts_on': start.isoformat(), 'ends_on': last.isoformat(), 'generated_at': now.isoformat(),
                'course_name': course_name, 'group_name': group_name, 'options': options, 'sessions': sessions,
                'students': student_rows, 'courses': course_rows, 'rows': rows, 'totals': summarize(rows)}

    def export_attendance_report(self, path, report):
        write_report_workbook(path, report)


def report_context(report):
    return (f"正式上课日期：{report['starts_on']} 至 {report['ends_on']}；课程：{report['course_name'] or '全部'}；"
            f"学生快照班级：{report['group_name'] or '全部'}；快照生成：{report['generated_at'].replace('T', ' ')}；只计已结束课次")


def summary_values(row):
    return [row[key] if row[key] is not None else '—' for key in SUMMARY_FIELDS]


def write_report_workbook(path, report):
    from classroom_attendance.maintenance.maintenance_tools import require_disk_space
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('统计报表必须使用 .xlsx 扩展名。')
    require_disk_space(path.parent)
    if (len(report['sessions']) + 3 > 16384 or len(report['rows']) > 1048572
            or max(len(report['students']), len(report['courses'])) > 1048571):
        raise ValueError('报表超过 Excel 行列上限，请缩小日期范围或筛选课程、班级。')
    workbook = Workbook()
    workbook.remove(workbook.active)
    temporary = None

    def sheet(title, headers, values, *, rate_columns=(), freeze='D5'):
        page = workbook.create_sheet(title)
        page.append([title])
        page.append([report_context(report)])
        page.append([])
        page.append(headers)
        for values_row in values:
            checkpoint()
            page.append(values_row)
        page.freeze_panes = freeze
        page.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{page.max_row}'
        page.sheet_view.showGridLines = False
        page['A1'].font = Font(name='Microsoft YaHei', size=16, bold=True, color='17243C')
        page.row_dimensions[1].height = 32
        page.row_dimensions[2].height = 42
        page['A2'].alignment = Alignment(wrap_text=True, vertical='center')
        page.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        page.row_dimensions[4].height = 60 if title == '考勤矩阵' else 42
        for cells in page:
            checkpoint()
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if cell.row == 4:
                    cell.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='3267E3')
                    cell.alignment = Alignment(wrap_text=True, vertical='center')
                elif cell.row >= 5:
                    cell.alignment = Alignment(vertical='center', wrap_text=True)
                    if cell.column in rate_columns:
                        cell.number_format = '0.00%'
        for column in range(1, len(headers) + 1):
            page.column_dimensions[get_column_letter(column)].width = 28 if column == 1 else 22
        return page

    try:
        totals = report['totals']
        student_values = [[row['student_id'], row['name'], row['group_name'], *summary_values(row)] for row in report['students']]
        student_values.append(['合计', '', '', *summary_values(totals)])
        sheet('学生学期汇总', ['学号', '姓名（最近快照）', '班级（最近快照）', *SUMMARY_HEADERS], student_values,
              rate_columns=(11, 12))
        course_values = [[row['course_name'], row['group_name'], len(row['session_ids']), *summary_values(row)] for row in report['courses']]
        course_values.append(['合计', '', len(report['sessions']), *summary_values(totals)])
        sheet('课程汇总', ['课程', '课次班级', '已结束课次数', *SUMMARY_HEADERS], course_values, rate_columns=(11, 12))
        headers = ['学号', '姓名（最近快照）', '班级（最近快照）'] + [
            f"#{session['id']} {session['course_name']}\n{session['group_name']}\n{session['starts_at'].replace('T', ' ')}"
            for session in report['sessions']]
        sheet('考勤矩阵', headers, ([row['student_id'], row['name'], row['group_name']] + [
            REPORT_STATUSES.get(row['cells'].get(session['id']), '—') for session in report['sessions']]
            for row in report['students']))
        sheet('考勤明细', ['课次编号', '课程', '课次班级', '正式上课时间', '学号', '姓名快照', '班级快照', '状态',
                         '实际到课时间', '到课时间说明', '记录方式', '备注', '首次人脸签到时间', '开放签到时间', '迟到起点', '关闭签到时间'],
              ([row['session_id'], row['course_name'], row['session_group_name'], row['session_starts_at'].replace('T', ' '),
                row['student_id'], row['name'], row['group_name'], REPORT_STATUSES[row['status']],
                (row['signed_at'] or '').replace('T', ' '),
                arrival_time_note(row), method_text(row['method'], short=True), row['note'],
                (row['original_signed_at'] or '').replace('T', ' '),
                *[row[key].replace('T', ' ') for key in ('session_opens_at', 'session_late_at', 'session_ends_at')]] for row in report['rows']), freeze='F5')
        rules = workbook.create_sheet('统计口径')
        for content in ['统计口径', report_context(report), *REPORT_RULES]:
            rules.append([content])
            rules.cell(rules.max_row, 1).data_type = 's'
            rules.cell(rules.max_row, 1).alignment = Alignment(wrap_text=True, vertical='center')
            rules.row_dimensions[rules.max_row].height = 44
        rules.column_dimensions['A'].width = 120
        rules.freeze_panes = 'A3'
        checkpoint()
        with tempfile.NamedTemporaryFile(prefix=f'.{path.stem}-', suffix='.tmp.xlsx', dir=path.parent, delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        checkpoint()
        temporary.replace(path)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
